"""
APCOTEX Excel parser — template-driven sales report extraction.

Architecture (after):
  Workbook → Official Template Detection → Distributor Labels → Sales Headers
           → Official mapping (bypass aliases) OR alias fallback
           → Row loop until one empty row → validate per-row → import valid

Architecture (before):
  Alias guess → silent empty result on mismatch → 0-row report
"""

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import pandas as pd
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from app.core.logging import get_logger
from app.exceptions import ExcelProcessingError
from app.integrations.excel.confidence import compute_confidence_breakdown
from app.integrations.excel.headers import (
    OFFICIAL_TEMPLATE_NAME,
    find_sales_header_row,
    is_distributor_label,
    match_distributor_field,
    normalize_header,
    resolve_sales_column_mapping,
    score_header_row,
)
from app.schemas.sales_record import ParsedSalesRow
from app.utils.hashing import build_sales_row_hash
from app.utils.quantity import optional_int, parse_quantity, safe_str
from app.utils.reporting_month import normalize_reporting_month
from app.validators.excel_validators import validate_customer_master_dataframe
from app.integrations.excel.product_master_parser import parse_product_master_workbook

logger = get_logger(__name__)


class SalesParseResult(BaseModel):
    """Sales parse output with extraction-quality metadata."""

    rows: List[ParsedSalesRow] = Field(default_factory=list)
    expected_rows: int = 0
    imported_rows: int = 0
    incomplete_rows: int = 0
    template_detected: bool = False
    sales_table_detected: bool = False
    template_name: str = "none"
    mapping_strategy: str = "none"
    distributor_details: Dict[str, str] = Field(default_factory=dict)
    reporting_month: Optional[str] = None
    quality_score: int = 0
    confidence_breakdown: Dict[str, Any] = Field(default_factory=dict)
    row_errors: List[str] = Field(default_factory=list)


class ExcelParserService:
    """Template-driven parser for APCOTEX sales and master-data workbooks."""

    def read_dataframe(self, path: Union[str, Path], sheet_name: Union[str, int] = 0) -> pd.DataFrame:
        """Load an Excel file into a DataFrame."""
        file_path = Path(path)
        if not file_path.exists():
            raise ExcelProcessingError(f"Excel file not found: {file_path}")
        try:
            logger.info("Workbook Loaded | path={}", file_path)
            engine = "xlrd" if file_path.suffix.lower() == ".xls" else "openpyxl"
            df = pd.read_excel(file_path, sheet_name=sheet_name, engine=engine)
        except Exception as exc:
            logger.exception("Failed to read Excel | path={}", file_path)
            raise ExcelProcessingError(f"Unable to read Excel file: {exc}") from exc
        df = df.dropna(how="all")
        df.columns = [str(c).strip() for c in df.columns]
        return df

    def parse_sales_report(self, path: Union[str, Path]) -> SalesParseResult:
        """
        Parse a distributor sales Excel.

        Priority:
          1. Official APCOTEX template (exact header contract)
          2. Alias-based fallback for legacy / variant headers

        Fail-fast on header mismatch. Never returns an empty success result.
        """
        file_path = Path(path)
        logger.info("Workbook Received | path={}", file_path)

        raw_rows, sheet_name = self._read_raw_rows_with_meta(file_path)
        if not raw_rows:
            raise ExcelProcessingError("Sales Excel is empty")

        logger.info(
            "Workbook Loaded | sheet={} | total_rows={} | total_cols={}",
            sheet_name,
            len(raw_rows),
            max((len(r) for r in raw_rows), default=0),
        )

        distributor_details = self._extract_distributor_details(raw_rows)
        logger.info(
            "Distributor Parsed | name={!r} | company={!r} | reporting_month={!r}",
            distributor_details.get("name") or "",
            distributor_details.get("company") or "",
            distributor_details.get("reporting_month") or "",
        )

        header_row_idx = find_sales_header_row(raw_rows)
        if header_row_idx is None:
            # Flat workbook fallback (first row = headers via pandas)
            table_df = self.read_dataframe(file_path)
            sales_table_detected = not table_df.empty
            logger.info(
                "Template Detection | status=legacy_flat_attempt | header_row=none | columns={}",
                list(table_df.columns),
            )
        else:
            headers = [
                safe_str(c) or f"col_{i}" for i, c in enumerate(raw_rows[header_row_idx])
            ]
            # Drop trailing empty header placeholders
            while headers and headers[-1].startswith("col_"):
                headers.pop()

            data_rows: List[List[Any]] = []
            for row in raw_rows[header_row_idx + 1 :]:
                padded = list(row[: len(headers)]) + [None] * max(0, len(headers) - len(row))
                padded = padded[: len(headers)]
                if not any(safe_str(c) for c in padded):
                    break
                data_rows.append(padded)

            if not data_rows:
                raise ExcelProcessingError(
                    "Official APCOTEX Template validation failed. "
                    "Sales table header found but no data rows.",
                    details={"header_row": header_row_idx + 1, "headers": headers},
                )

            table_df = pd.DataFrame(data_rows, columns=headers)
            table_df.columns = [str(c).strip() for c in table_df.columns]
            sales_table_detected = True
            logger.info(
                "Header Found | excel_row={} | columns={} | normalized={}",
                header_row_idx + 1,
                list(table_df.columns),
                [normalize_header(c) for c in table_df.columns],
            )

        # --- Column mapping (official first, alias fallback) ---
        try:
            mapping, strategy, is_official = resolve_sales_column_mapping(list(table_df.columns))
        except ExcelProcessingError:
            logger.error(
                "Column Mapping FAILED | columns={} | normalized={}",
                list(table_df.columns),
                [normalize_header(c) for c in table_df.columns],
            )
            raise

        template_name = OFFICIAL_TEMPLATE_NAME if is_official else "Legacy / Alias Fallback"
        template_detected = is_official or bool(
            distributor_details.get("name")
            or distributor_details.get("company")
        )

        logger.info(
            "Template Detection | template={!r} | status={} | strategy={}",
            template_name,
            "Matched" if is_official else "Fallback",
            strategy,
        )
        logger.info(
            "Column Mapping | strategy={} | mapping={}",
            strategy,
            {k: mapping[k] for k in mapping},
        )

        header_distributor = distributor_details.get("name") or ""
        if "distributor" not in mapping and not header_distributor:
            raise ExcelProcessingError(
                "Official APCOTEX Template validation failed. "
                "Missing Name of Person in report metadata "
                "and no Distributor column in the sales table.",
                details={"distributor_details": distributor_details},
            )

        # --- Reporting Quarter (stored as reporting_month key; required) ---
        reporting_month = normalize_reporting_month(
            distributor_details.get("reporting_month")
        )
        if not reporting_month and "period" in mapping:
            for _, series in table_df.iterrows():
                candidate = normalize_reporting_month(series.get(mapping["period"]))
                if candidate:
                    reporting_month = candidate
                    break
        if not reporting_month:
            raise ExcelProcessingError(
                "Official APCOTEX Template validation failed. "
                "Missing Reporting Quarter in report metadata.",
                details={"distributor_details": distributor_details},
            )
        distributor_details["reporting_month"] = reporting_month
        # Company defaults to Name of Person when Company Name row is absent
        if not (distributor_details.get("company") or "").strip():
            distributor_details["company"] = header_distributor
        logger.info("Reporting Quarter Resolved | value={!r}", reporting_month)

        # --- Row processing ---
        rows: List[ParsedSalesRow] = []
        errors: List[str] = []
        incomplete_rows = 0
        expected_rows = 0
        default_distributor = header_distributor
        default_company = distributor_details.get("company") or header_distributor

        for index, series in table_df.iterrows():
            if self._is_completely_empty_row(series, mapping):
                break

            expected_rows += 1
            excel_row_num = (
                (header_row_idx + 2 + int(index)) if header_row_idx is not None else int(index) + 2
            )
            row_errors: List[str] = []

            try:
                if "distributor" in mapping:
                    distributor = safe_str(series.get(mapping["distributor"])) or default_distributor
                else:
                    distributor = default_distributor

                customer_name = safe_str(series.get(mapping["customer_name"]))
                segment = safe_str(series.get(mapping["segment"]))
                product = safe_str(series.get(mapping["product"]))
                quantity_raw = series.get(mapping["quantity"])

                row_period = reporting_month
                if "period" in mapping:
                    legacy = safe_str(series.get(mapping["period"]))
                    if legacy:
                        row_period = legacy

                if not distributor:
                    row_errors.append("Missing Name of Person / Distributor")
                if not customer_name:
                    row_errors.append("Missing Customer Name")
                if not segment:
                    row_errors.append("Missing Segment")
                if not product:
                    row_errors.append("Missing Product")

                try:
                    quantity, quantity_display = parse_quantity(quantity_raw)
                    if quantity <= 0:
                        row_errors.append("Sales Quantity must be a positive number")
                        quantity, quantity_display = None, ""  # type: ignore[assignment]
                except ValueError as qty_err:
                    row_errors.append(f"Invalid Sales Quantity ({qty_err})")
                    quantity, quantity_display = None, ""  # type: ignore[assignment]

                sr_no = None
                if "sr_no" in mapping:
                    sr_no = optional_int(series.get(mapping["sr_no"]))
                if sr_no is None:
                    sr_no = expected_rows

                if row_errors or quantity is None:
                    incomplete_rows += 1
                    reason = "; ".join(row_errors) if row_errors else "Invalid Sales Quantity"
                    msg = f"Row {sr_no} (excel row {excel_row_num}): {reason}"
                    errors.append(msg)
                    logger.warning("Row Skipped | {}", msg)
                    continue

                row_hash = build_sales_row_hash(
                    distributor, customer_name, segment, product, quantity, row_period
                )
                rows.append(
                    ParsedSalesRow(
                        sr_no=sr_no,
                        distributor=distributor,
                        customer_name=customer_name,
                        segment=segment,
                        product=product,
                        quantity=quantity,
                        quantity_display=quantity_display,
                        period=row_period,
                        company=default_company or None,
                        row_hash=row_hash,
                        errors=[],
                    )
                )
            except Exception as exc:
                incomplete_rows += 1
                msg = f"Row {excel_row_num}: {exc}"
                errors.append(msg)
                logger.warning("Row Skipped | {}", msg)

        imported_rows = len(rows)
        logger.info(
            "Rows Read={} | Rows Imported={} | Rows Skipped={} | Validation Errors={}",
            expected_rows,
            imported_rows,
            incomplete_rows,
            len(errors),
        )
        if errors:
            for err in errors[:20]:
                logger.info("Validation Error | {}", err)

        if imported_rows == 0:
            breakdown = compute_confidence_breakdown(
                template_detected=template_detected,
                sales_table_detected=sales_table_detected,
                is_official_template=is_official,
                distributor_details=distributor_details,
                expected_rows=expected_rows,
                imported_rows=0,
                incomplete_rows=incomplete_rows,
                parse_succeeded=False,
                mapping_strategy=strategy,
            )
            quality = int(breakdown["total"])
            raise ExcelProcessingError(
                "No valid sales rows found in Excel",
                details={
                    "errors": errors,
                    "expected_rows": expected_rows,
                    "imported_rows": 0,
                    "quality_score": quality,
                    "confidence_breakdown": breakdown,
                    "template": template_name,
                    "strategy": strategy,
                },
            )

        breakdown = compute_confidence_breakdown(
            template_detected=template_detected,
            sales_table_detected=sales_table_detected,
            is_official_template=is_official,
            distributor_details=distributor_details,
            expected_rows=expected_rows,
            imported_rows=imported_rows,
            incomplete_rows=incomplete_rows,
            parse_succeeded=True,
            mapping_strategy=strategy,
        )
        quality = int(breakdown["total"])

        logger.info(
            "Confidence Score={} | breakdown={} | template={!r} | distributor={!r} | reporting_month={!r}",
            quality,
            breakdown.get("components"),
            template_name,
            default_distributor or (rows[0].distributor if rows else ""),
            reporting_month,
        )
        if expected_rows != imported_rows:
            logger.warning(
                "Row mismatch | Expected={} | Imported={} | delta={}",
                expected_rows,
                imported_rows,
                expected_rows - imported_rows,
            )

        return SalesParseResult(
            rows=rows,
            expected_rows=expected_rows,
            imported_rows=imported_rows,
            incomplete_rows=incomplete_rows,
            template_detected=template_detected,
            sales_table_detected=sales_table_detected,
            template_name=template_name,
            mapping_strategy=strategy,
            distributor_details=distributor_details,
            reporting_month=reporting_month,
            quality_score=quality,
            confidence_breakdown=breakdown,
            row_errors=errors,
        )

    @staticmethod
    def _is_completely_empty_row(series: pd.Series, mapping: Dict[str, str]) -> bool:
        """True when all mapped sales fields are blank (stop condition)."""
        fields = ["customer_name", "segment", "product", "quantity", "sr_no"]
        if "distributor" in mapping:
            fields.append("distributor")
        for field in fields:
            col = mapping.get(field)
            if col is None:
                continue
            if safe_str(series.get(col)):
                return False
            raw = series.get(col)
            if field == "quantity" and raw is not None and str(raw).strip() not in {"", "nan", "None"}:
                return False
        return True

    def _read_raw_rows_with_meta(self, file_path: Path) -> Tuple[List[List[Any]], str]:
        """Read sheet as a grid; return (rows, sheet_name)."""
        suffix = file_path.suffix.lower()
        try:
            if suffix == ".xls":
                engine_df = pd.read_excel(file_path, sheet_name=0, header=None, engine="xlrd")
                rows = [
                    [None if pd.isna(v) else v for v in row]
                    for row in engine_df.values.tolist()
                ]
                return rows, "Sheet1"
            workbook = load_workbook(file_path, data_only=True, read_only=False)
            sheet = workbook.active
            sheet_name = sheet.title
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            workbook.close()
            return rows, sheet_name
        except Exception as exc:
            logger.exception("Failed to read workbook grid | path={}", file_path)
            raise ExcelProcessingError(f"Unable to read Excel file: {exc}") from exc

    def _extract_distributor_details(self, raw_rows: List[List[Any]]) -> Dict[str, str]:
        """
        Extract Distributor Details by label (official contract).

        Phone is never merged into Address — phone labels terminate address
        continuation and are stored separately.
        """
        details: Dict[str, str] = {
            "name": "",
            "company": "",
            "address": "",
            "phone": "",
            "reporting_month": "",
        }
        scan_limit = min(len(raw_rows), 50)
        i = 0
        while i < scan_limit:
            row = raw_rows[i]
            cells = [c for c in row if c is not None and str(c).strip() != ""]
            if not cells:
                i += 1
                continue

            if score_header_row(row) >= 4:
                break

            label_raw: Any = cells[0]
            raw_value: Any = None
            value = ""
            if len(cells) >= 2:
                raw_value = cells[1]
                value = safe_str(cells[1])
            elif len(cells) == 1 and ":" in str(cells[0]):
                left, _, right = str(cells[0]).partition(":")
                label_raw = left
                raw_value = right
                value = safe_str(right)

            field = match_distributor_field(label_raw)
            if not field:
                i += 1
                continue

            if field == "address":
                # Capture same-row value plus continuation lines until next label
                parts: List[str] = []
                if value:
                    parts.append(value)
                j = i + 1
                while j < scan_limit:
                    cont = raw_rows[j]
                    cont_cells = [c for c in cont if c is not None and str(c).strip() != ""]
                    if not cont_cells:
                        break
                    if score_header_row(cont) >= 4:
                        break
                    # Phone / Name / Company labels must not fold into address
                    if is_distributor_label(cont_cells[0]):
                        break
                    if len(cont_cells) >= 2 and is_distributor_label(cont_cells[0]):
                        break
                    parts.append("\n".join(safe_str(c) for c in cont_cells if safe_str(c)))
                    j += 1
                if parts and not details["address"]:
                    details["address"] = "\n".join(parts).strip()
                i = j
                continue

            if field == "reporting_month":
                month = normalize_reporting_month(raw_value if raw_value is not None else value)
                if month and not details[field]:
                    details[field] = month
            elif value and not details[field]:
                details[field] = value
            i += 1

        return details

    def parse_customer_master(self, path: Union[str, Path]):
        """
        Parse customer master Excel (Phase 2).

        Extracts CUSTOMER NAME only: trim, drop blanks, de-duplicate (case-insensitive).
        All other columns are ignored.

        Returns ``CustomerMasterParseResult`` with explicit skip reasons
        (blank name, duplicate, validation). Fully blank spreadsheet rows
        are counted as Blank Customer Name — never silent.
        """
        import pandas as pd

        from app.schemas.master_parse import CustomerMasterParseResult

        file_path = Path(path)
        if not file_path.exists():
            raise ExcelProcessingError(f"Excel file not found: {file_path}")
        try:
            engine = "xlrd" if file_path.suffix.lower() == ".xls" else "openpyxl"
            df_raw = pd.read_excel(file_path, sheet_name=0, engine=engine)
        except Exception as exc:
            raise ExcelProcessingError(f"Unable to read Excel file: {exc}") from exc

        df_raw.columns = [str(c).strip() for c in df_raw.columns]
        excel_rows = int(len(df_raw))
        df = df_raw.dropna(how="all")
        fully_blank_rows = excel_rows - int(len(df))

        mapping = validate_customer_master_dataframe(df if not df.empty else df_raw)
        name_col = mapping["customer_name"]

        seen: set[str] = set()
        records: List[dict] = []
        blank_customer_name = fully_blank_rows
        duplicate_names = 0
        validation_errors = 0
        max_len = 255

        for raw in df[name_col].tolist():
            name = safe_str(raw)
            if not name:
                blank_customer_name += 1
                continue
            if len(name) > max_len:
                validation_errors += 1
                continue
            key = name.casefold()
            if key in seen:
                duplicate_names += 1
                continue
            seen.add(key)
            records.append({"customer_name": name})

        if not records:
            raise ExcelProcessingError("No valid customer names found in Customer Master")

        result = CustomerMasterParseResult(
            records=records,
            excel_rows=excel_rows,
            blank_customer_name=blank_customer_name,
            duplicate_names=duplicate_names,
            validation_errors=validation_errors,
        )
        logger.info(
            "Parsed customer master | excel_rows={} imported={} blank={} "
            "duplicates={} validation_errors={} skipped={}",
            result.excel_rows,
            result.imported,
            result.blank_customer_name,
            result.duplicate_names,
            result.validation_errors,
            result.skipped,
        )
        return result

    def parse_product_master(self, path: Union[str, Path]) -> tuple[List[dict], int]:
        """
        Parse Product Master Excel (production).

        Supports:
          - Single continuous table (column A+)
          - Multi-block side-by-side APCOTEX layouts (B:C, E:F, H:I, K:L, …)
          - Robust header aliases / normalization

        Returns (records, duplicates_ignored).
        """
        result = parse_product_master_workbook(path)
        return result.records, result.duplicates_ignored
