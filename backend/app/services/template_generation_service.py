"""Official distributor template generation & download (quarterly workflow)."""

import shutil
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from sqlalchemy.orm import Session

from app.constants import DISTRIBUTOR_TEMPLATE_FILENAME
from app.core.config import settings
from app.core.logging import get_logger
from app.enums import AuditAction
from app.exceptions import GraphAPIError, ValidationAppError
from app.integrations.excel.template_generator import ExcelTemplateGenerator
from app.integrations.graph.client import GraphClient
from app.repositories.distributor_customer_mapping_repository import (
    DistributorCustomerMappingRepository,
)
from app.repositories.distributor_repository import DistributorRepository
from app.repositories.master_repository import ProductMasterRepository
from app.repositories.sales_record_repository import SalesRecordRepository
from app.schemas.audit import AuditTrailCreate
from app.schemas.dashboard import TemplateGenerateResponse
from app.schemas.distributor import EmailDraftResponse
from app.services.audit_service import AuditService
from app.services.email_package_service import sanitize_filename_part
from app.utils.files import ensure_dir
from app.utils.reporting_month import normalize_reporting_month

logger = get_logger(__name__)


class TemplateGenerationService:
    """Generate and serve the official APCOTEX quarterly distributor template."""

    def __init__(self, db: Session) -> None:
        self.db = db
        self.products = ProductMasterRepository(db)
        self.sales = SalesRecordRepository(db)
        self.distributors = DistributorRepository(db)
        self.customer_maps = DistributorCustomerMappingRepository(db)
        self.templates = ExcelTemplateGenerator()
        self.graph = GraphClient()
        self.audit = AuditService(db)

    def _template_path(self) -> Path:
        return Path(settings.download_dir) / DISTRIBUTOR_TEMPLATE_FILENAME

    def _version(self) -> str:
        now = datetime.now(timezone.utc)
        return f"v{now.year}.{now.month:02d}.{now.day:02d}"

    def generate(
        self,
        *,
        actor: str = "system",
        mode: str = "generic",
        distributor_id: Optional[int] = None,
        reporting_quarter: Optional[str] = None,
        output_path: Optional[Path] = None,
        update_canonical: bool = True,
    ) -> TemplateGenerateResponse:
        """
        Build quarterly template from Product Master.

        ``generic``      → Customer Name free text
        ``distributor``  → Customer Name dropdown from all historical customers
                           for the selected distributor (Reporting Quarter is
                           entered later in Excel by the distributor)
        """
        started = time.perf_counter()
        mode_key = (mode or "generic").strip().lower()
        distributor_mode = mode_key in {"distributor", "distributor_specific", "specific"}

        products_by_segment = self.products.list_codes_by_segment()
        product_count = sum(len(v) for v in products_by_segment.values())
        if not products_by_segment:
            raise ValidationAppError(
                "Cannot generate template: Product Master is empty. Upload Product Master first."
            )

        customers: list[str] = []
        contact_person: Optional[str] = None
        company_name: Optional[str] = None
        dist = None
        warning: Optional[str] = None
        fallback_generic = False
        effective_distributor_mode = distributor_mode

        # Optional prefill only (Outlook package / advanced callers). Not required.
        quarter = None
        if reporting_quarter:
            quarter = normalize_reporting_month(reporting_quarter) or reporting_quarter.strip()

        if distributor_mode:
            if not distributor_id:
                raise ValidationAppError(
                    "distributor_id is required for Distributor-Specific Template"
                )
            dist = self.distributors.get_or_raise(distributor_id)
            if not dist.is_active or dist.is_deleted:
                raise ValidationAppError("Distributor is inactive or deleted")
            customers = self.customer_maps.list_names_for_template(
                distributor_id, quarter
            )
            contact_person = (dist.contact_person or dist.name or "").strip() or None
            company_name = (dist.company or dist.name or "").strip() or None
            if not customers:
                effective_distributor_mode = False
                fallback_generic = True
                warning = (
                    "No customer history found for this distributor. "
                    "A generic template has been generated."
                )
                logger.info(
                    "Distributor template fallback to generic | distributor_id={} | quarter={}",
                    distributor_id,
                    quarter,
                )

        version = self._version()
        if distributor_mode and dist is not None:
            label = sanitize_filename_part(dist.company or dist.name)
            q_part = sanitize_filename_part(quarter or "Template")
            file_name = f"Apcotex_{label}_{q_part}_{version}.xlsx"
        else:
            file_name = f"Apcotex_Distributor_Template_{version}.xlsx"

        output = output_path or (Path(settings.download_dir) / file_name)
        ensure_dir(output.parent)

        gen_started = time.perf_counter()
        path = self.templates.generate(
            products_by_segment=products_by_segment,
            customers=customers if effective_distributor_mode else None,
            mode="distributor" if effective_distributor_mode else "generic",
            contact_person=contact_person,
            company_name=company_name,
            reporting_quarter=quarter,
            output_path=output,
        )
        gen_ms = round((time.perf_counter() - gen_started) * 1000, 2)

        if update_canonical and not distributor_mode:
            canonical = self._template_path()
            if path.resolve() != canonical.resolve():
                shutil.copyfile(path, canonical)

        generated_at = datetime.now(timezone.utc).isoformat()
        elapsed_ms = round((time.perf_counter() - started) * 1000, 2)
        if fallback_generic:
            mode_label = "generic-fallback"
        elif effective_distributor_mode:
            mode_label = "distributor-specific"
        else:
            mode_label = "generic"
        self.audit.log(
            AuditTrailCreate(
                user_name=actor,
                action=AuditAction.GENERATED,
                details=(
                    f"Generated {mode_label} quarterly template {file_name} "
                    f"({product_count:,} products, {len(customers)} customers)"
                    + (f" | {warning}" if warning else "")
                ),
                entity_type="template",
                module="Master Data",
                status="Success",
                extra_metadata={
                    "file_name": file_name,
                    "mode": mode_label,
                    "distributor_id": distributor_id,
                    "reporting_quarter": quarter,
                    "products": product_count,
                    "customers": len(customers),
                    "fallback_generic": fallback_generic,
                    "elapsed_ms": elapsed_ms,
                },
            )
        )
        logger.info(
            "Template generated | file={} mode={} products={} customers={} gen_ms={} total_ms={}",
            file_name,
            mode_label,
            product_count,
            len(customers),
            gen_ms,
            elapsed_ms,
        )
        if fallback_generic:
            message = warning or "A generic template has been generated."
        elif effective_distributor_mode:
            message = (
                "Distributor-specific template generated successfully "
                f"({len(customers)} historical customers)"
            )
        else:
            message = "Distributor template generated successfully"
        return TemplateGenerateResponse(
            success=True,
            status="ready",
            message=message,
            template_version=version,
            file_name=file_name,
            customers_count=len(customers),
            products_count=product_count,
            generated_at=generated_at,
            mode=mode_label,
            warning=warning,
            fallback_generic=fallback_generic,
        )

    def resolve_download_path(self, *, preferred_name: Optional[str] = None) -> tuple[Path, str]:
        """Return (path, download_filename) for the latest generated template."""
        download_dir = Path(settings.download_dir)
        if preferred_name:
            # Prevent path traversal — only a bare filename under download_dir
            safe_name = Path(str(preferred_name)).name
            candidate = (download_dir / safe_name).resolve()
            if candidate.is_file() and candidate.parent == download_dir.resolve():
                return candidate, safe_name

        # Prefer newest distributor-specific or generic template under downloads/
        matches = sorted(
            [
                p
                for p in download_dir.glob("Apcotex_*.xlsx")
                if p.is_file() and not p.name.startswith("~$")
            ],
            key=lambda p: p.stat().st_mtime,
            reverse=True,
        )
        if matches:
            return matches[0], matches[0].name

        canonical = self._template_path()
        if canonical.is_file():
            return canonical, DISTRIBUTOR_TEMPLATE_FILENAME

        raise ValidationAppError(
            "No generated template found. Call POST /template/generate first."
        )

    def generate_quarterly_package(
        self,
        *,
        distributor_id: int,
        reporting_quarter: str,
        actor: str = "system",
    ) -> EmailDraftResponse:
        """Backward-compatible alias → create Microsoft Graph Outlook draft."""
        return self.create_email_draft(
            distributor_id=distributor_id,
            reporting_quarter=reporting_quarter,
            actor=actor,
        )

    def create_email_draft(
        self,
        *,
        distributor_id: int,
        reporting_quarter: str,
        actor: str = "system",
    ) -> EmailDraftResponse:
        """
        Generate distributor-specific Excel (same TemplateGenerationService path)
        and create a Microsoft Graph **draft** in GRAPH_MAILBOX (not sent).
        """
        dist = self.distributors.get_or_raise(distributor_id)
        if not dist.is_active or dist.is_deleted:
            raise ValidationAppError("Distributor is inactive or deleted")

        to_email = (dist.email or "").strip()
        if not to_email:
            raise ValidationAppError(
                "Distributor email address is not configured. "
                "Please update the distributor details before creating the email draft."
            )

        quarter = normalize_reporting_month(reporting_quarter) or reporting_quarter.strip()
        if not quarter:
            raise ValidationAppError("reporting_quarter is required (e.g. FY 2025-26 • Q2)")

        company_label = (dist.company or dist.name or "Distributor").strip()
        label = sanitize_filename_part(company_label)
        q_part = sanitize_filename_part(quarter)
        excel_name = f"Apcotex_{label}_{q_part}_Template.xlsx"
        packages_dir = Path(settings.download_dir) / "packages"
        ensure_dir(packages_dir)
        excel_path = packages_dir / excel_name

        if excel_path.exists():
            excel_path.unlink()

        try:
            result = self.generate(
                actor=actor,
                mode="distributor",
                distributor_id=distributor_id,
                reporting_quarter=quarter,
                output_path=excel_path,
                update_canonical=False,
            )
        except ValidationAppError as exc:
            msg = str(exc.message if hasattr(exc, "message") else exc)
            if "Product Master" in msg or "empty" in msg.lower():
                raise ValidationAppError(
                    "Product Master is not available. Please upload the Product Master "
                    "before creating the distributor email draft."
                ) from exc
            raise

        mapped_count = self.customer_maps.count_active(distributor_id)
        if mapped_count > 0 and (result.fallback_generic or result.customers_count <= 0):
            raise ValidationAppError(
                "Distributor has mapped customers but the email draft template was generated "
                "without a Customer Name dropdown. Re-check customer mappings / sales history.",
                details={
                    "distributor_id": distributor_id,
                    "mapped_count": mapped_count,
                    "customers_count": result.customers_count,
                    "mode": result.mode,
                },
            )

        if not excel_path.is_file():
            raise ValidationAppError(f"Draft Excel was not written to {excel_path.name}")

        self._assert_customer_dropdown(
            excel_path, expect_dropdown=result.customers_count > 0
        )

        contact = (dist.contact_person or "").strip()
        greeting = contact if contact else "Team"
        subject = f"Distributor Sales Template – {quarter} – {company_label}"
        body = (
            f"Dear {greeting},\n\n"
            f"Please find attached the Distributor Sales Template for {quarter}.\n\n"
            "Kindly complete the required sales information and share the completed "
            "template with the APCOTEX team as per the agreed process.\n\n"
            "Please ensure that all required fields are completed before submitting "
            "the template.\n\n"
            "Regards,\n"
            "APCOTEX Team\n"
        )
        cc_email = (dist.cc_email or "").strip() or None
        excel_bytes = excel_path.read_bytes()

        try:
            mailbox = self.graph.resolve_mailbox()
            draft = self.graph.create_draft_message(
                subject=subject,
                body_text=body,
                to_email=to_email,
                cc_email=cc_email,
                mailbox=mailbox,
            )
            draft_id = str(draft.get("id") or "")
            self.graph.add_file_attachment(
                draft_id,
                filename=excel_name,
                content=excel_bytes,
                mailbox=mailbox,
            )
        except ValidationAppError:
            raise
        except GraphAPIError as exc:
            logger.exception(
                "Graph draft creation failed | distributor_id={} | quarter={}",
                distributor_id,
                quarter,
            )
            details = getattr(exc, "details", None)
            detail_text = str(details or "")
            if "403" in detail_text or "Authorization_RequestDenied" in detail_text:
                logger.error(
                    "Microsoft Graph Mail.ReadWrite permission/admin consent is required "
                    "to create drafts in {}",
                    settings.graph_mailbox or "(GRAPH_MAILBOX)",
                )
            raise GraphAPIError(
                "Unable to create the Outlook draft. Please contact the administrator. "
                "Microsoft Graph Mail.ReadWrite permission/admin consent may be required.",
                details={"safe": True},
            ) from exc
        except Exception as exc:  # noqa: BLE001
            logger.exception("Unexpected failure creating Outlook draft")
            raise GraphAPIError(
                "Unable to create the Outlook draft. Please contact the administrator."
            ) from exc

        self.audit.log(
            AuditTrailCreate(
                user_name=actor,
                action=AuditAction.GENERATED,
                details=(
                    f"Created Distributor Email Draft | distributor={company_label} | "
                    f"quarter={quarter} | recipient={to_email} | "
                    f"attachment={excel_name} | draft_id={draft_id}"
                ),
                entity_type="distributor",
                entity_id=str(distributor_id),
                module="Distributor Communication",
                status="Success",
                extra_metadata={
                    "draft_id": draft_id,
                    "mailbox": mailbox,
                    "recipient": to_email,
                    "cc": cc_email,
                    "attachment_name": excel_name,
                    "reporting_quarter": quarter,
                    "customers_count": result.customers_count,
                    "mode": result.mode,
                },
            )
        )
        logger.info(
            "Email draft created | mailbox={} | draft_id={} | distributor_id={} | attachment={}",
            mailbox,
            draft_id,
            distributor_id,
            excel_name,
        )
        return EmailDraftResponse(
            success=True,
            message=(
                f"Email draft created successfully in {mailbox}. "
                "Please open Outlook → Drafts to review and send."
            ),
            mailbox=mailbox,
            distributor_id=distributor_id,
            distributor_name=company_label,
            reporting_quarter=quarter,
            draft_id=draft_id,
            attachment_name=excel_name,
            recipient=to_email,
            cc=cc_email,
        )

    @staticmethod
    def _assert_customer_dropdown(excel_path: Path, *, expect_dropdown: bool) -> None:
        """Fail fast if the packaged workbook is missing Customer Name validation."""
        from openpyxl import load_workbook

        wb = load_workbook(excel_path, read_only=False, data_only=False)
        try:
            has_name = "DistributorCustomers" in wb.defined_names
            sheet = wb[wb.sheetnames[0]]
            formulas = [
                str(dv.formula1 or "") for dv in sheet.data_validations.dataValidation
            ]
            has_dv = any("DistributorCustomers" in f for f in formulas)
            if expect_dropdown and not (has_name and has_dv):
                raise ValidationAppError(
                    "Generated package Excel is missing Customer Name dropdown "
                    "(DistributorCustomers). Refusing to create Outlook draft.",
                    details={
                        "excel": excel_path.name,
                        "defined_names": list(wb.defined_names.keys()),
                        "formulas": formulas,
                    },
                )
        finally:
            wb.close()
