"""Bulk Persona and Employee-Distributor Mapping Import Service."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import bcrypt
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.core.logging import get_logger
from app.enums import AuditAction, UserRole
from app.exceptions import ValidationAppError
from app.models.distributor import Distributor
from app.models.user import User
from app.models.user_distributor import UserDistributor
from app.models.user_segment import UserSegment
from app.repositories.distributor_repository import DistributorRepository
from app.repositories.user_repository import UserRepository
from app.repositories.user_segment_repository import UserSegmentRepository
from app.schemas.audit import AuditTrailCreate
from app.services.audit_service import AuditService
from app.utils.distributor_name import normalize_company_name
from app.utils.files import ensure_dir

logger = get_logger(__name__)

TEMP_DEFAULT_PASSWORD = "Sales@123"

# Never create persona mappings for these distributors (not finalized).
_SKIP_DISTRIBUTOR_KEYWORDS = ("bajaj",)

# Excel owner labels that map onto the existing Admin (Deb) login — not a new sales user.
_ADMIN_OWNER_ALIASES = {
    "fyip dc",
    "fyip",
    "fyipdc",
    "dc",
    "deb",
    "debabrata",
    "debrabata",
    "debabrata c",
    "mr. debrabata",
    "mr debrabata",
    "mr. debabrata",
    "mr debabrata",
}


def _hash_pw(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


class BulkPersonaImportService:
    """Parses distributor mapping Excel and auto-populates users, segments, and distributor mappings."""

    def __init__(self, db: Session) -> None:
        self.db = db
        self.users = UserRepository(db)
        self.distributors = DistributorRepository(db)
        self.user_segments = UserSegmentRepository(db)
        self.audit = AuditService(db)

    @staticmethod
    def _clean_str(val: Any) -> str:
        if val is None:
            return ""
        return str(val).replace("\xa0", " ").strip()

    @staticmethod
    def _is_skipped_distributor(name: str) -> bool:
        key = (name or "").casefold()
        return any(token in key for token in _SKIP_DISTRIBUTOR_KEYWORDS)

    @staticmethod
    def _is_admin_owner(owner: str) -> bool:
        key = " ".join((owner or "").casefold().split())
        if key in _ADMIN_OWNER_ALIASES:
            return True
        compact = key.replace(" ", "").replace(".", "")
        if compact in {"fyipdc", "fyip", "debabrata", "debrabata"}:
            return True
        return "fyip" in key

    @classmethod
    def _generate_username(cls, full_name: str, existing_usernames: set[str]) -> str:
        """Generate first-name based lowercase username (e.g. 'Anup Pandey' -> 'anup')."""
        from app.services.auth_service import HARDCODED_SUPER_ADMIN_USERNAME

        parts = [p.strip().lower() for p in full_name.split() if p.strip()]
        base = parts[0] if parts else "user"
        base = "".join(c for c in base if c.isalnum()) or "user"
        candidate = base
        idx = 1
        reserved = existing_usernames | {HARDCODED_SUPER_ADMIN_USERNAME, "admin"}
        while candidate in reserved:
            candidate = f"{base}{idx}"
            idx += 1
        existing_usernames.add(candidate)
        return candidate

    def _resolve_admin_user(self) -> Optional[User]:
        """Prefer username=admin, else first active database admin."""
        admin = self.db.scalar(
            select(User).where(
                User.username == "admin",
                User.is_deleted.is_(False),
            )
        )
        if admin:
            return admin
        return self.db.scalar(
            select(User)
            .where(User.role == UserRole.ADMIN.value, User.is_deleted.is_(False))
            .order_by(User.id.asc())
        )

    def _load_master_distributors(self) -> List[Distributor]:
        return list(
            self.db.scalars(
                select(Distributor)
                .where(Distributor.is_deleted.is_(False))
                .order_by(Distributor.company.asc())
            ).all()
        )

    def _find_master_distributor(
        self, name: str, master: List[Distributor]
    ) -> Optional[Distributor]:
        """Match Excel Second Party to an existing distributor when possible."""
        from difflib import SequenceMatcher

        target = normalize_company_name(name)
        if not target:
            return None
        target_cf = target.casefold()

        for d in master:
            company = normalize_company_name(d.company or "").casefold()
            disp = normalize_company_name(d.name or "").casefold()
            if company == target_cf or disp == target_cf:
                return d

        candidates: List[Tuple[int, Distributor]] = []
        for d in master:
            company = normalize_company_name(d.company or "").casefold()
            disp = normalize_company_name(d.name or "").casefold()
            for hay in (company, disp):
                if not hay:
                    continue
                if target_cf in hay or hay in target_cf:
                    score = min(len(target_cf), len(hay))
                    candidates.append((score, d))
                    break
        if candidates:
            candidates.sort(key=lambda t: (-t[0], (t[1].company or t[1].name or "").casefold()))
            return candidates[0][1]

        fuzzy: List[Tuple[float, Distributor]] = []
        for d in master:
            for hay in (
                normalize_company_name(d.company or "").casefold(),
                normalize_company_name(d.name or "").casefold(),
            ):
                if not hay:
                    continue
                ratio = SequenceMatcher(None, target_cf, hay).ratio()
                if ratio >= 0.82:
                    fuzzy.append((ratio, d))
                    break
        if not fuzzy:
            return None
        fuzzy.sort(key=lambda t: -t[0])
        return fuzzy[0][1]

    def _resolve_or_create_distributor(
        self,
        name: str,
        *,
        owner_name: str,
        master: List[Distributor],
    ) -> Tuple[Distributor, bool]:
        """
        Resolve Second Party to a distributor row.

        Prefer existing Master match; otherwise create the company in Master.
        Bajaj is already filtered out at parse time.
        Returns ``(distributor, created)``.
        """
        existing = self._find_master_distributor(name, master)
        if existing:
            return existing, False

        created = self.distributors.get_or_create_by_company(
            company=name,
            representative_name=owner_name,
        )
        # Keep in-memory master list fresh for later rows
        if not any(d.id == created.id for d in master):
            master.append(created)
        return created, True

    def parse_excel(self, file_path: Union[str, Path]) -> List[Dict[str, str]]:
        """
        Parse Excel and extract rows for Second Party (Distributor), Owner (Sales Owner), Segment.
        Skips Bajaj rows. Keeps FYIP DC rows (mapped to Admin at execute time).
        """
        path = Path(file_path)
        try:
            wb = load_workbook(path, data_only=True)
        except Exception as exc:
            raise ValidationAppError(f"Unable to read Excel file: {exc}") from exc

        ws = wb.active
        if not ws:
            raise ValidationAppError("Uploaded Excel workbook is empty")

        matrix_rows = []
        for r in ws.iter_rows(values_only=True):
            if r and any(c is not None and str(c).strip() for c in r):
                matrix_rows.append(list(r))
        wb.close()

        if not matrix_rows:
            raise ValidationAppError("Uploaded Excel sheet contains no readable rows")

        header_idx: Optional[int] = None
        col_distributor: Optional[int] = None
        col_owner: Optional[int] = None
        col_segment: Optional[int] = None

        for idx, row in enumerate(matrix_rows[:25]):
            cells = [self._clean_str(c).lower() for c in row]
            c_dist = None
            c_own = None
            c_seg = None
            for col_i, cell_text in enumerate(cells):
                if not cell_text:
                    continue
                if cell_text in {
                    "second party",
                    "second party name",
                    "second party (distributor)",
                    "distributor",
                    "distributor name",
                    "distributor company",
                } and c_dist is None:
                    c_dist = col_i
                elif cell_text in {"owner", "sales owner", "owner name", "employee", "employee name"} and c_own is None:
                    c_own = col_i
                elif cell_text in {"segment", "product segment", "business segment"} and c_seg is None:
                    c_seg = col_i

            if c_dist is not None and c_own is not None:
                header_idx = idx
                col_distributor = c_dist
                col_owner = c_own
                col_segment = c_seg
                break

        if header_idx is None or col_distributor is None or col_owner is None:
            raise ValidationAppError(
                "Could not locate required columns in Excel. Ensure exact headers: 'Owner', 'Segment', and 'Second Party'."
            )

        parsed: List[Dict[str, str]] = []
        seen_keys: set[tuple[str, str]] = set()
        skipped_bajaj = 0

        for row in matrix_rows[header_idx + 1 :]:
            if len(row) <= max(col_distributor, col_owner):
                continue

            raw_dist = self._clean_str(row[col_distributor])
            raw_owner = self._clean_str(row[col_owner])
            raw_seg = (
                self._clean_str(row[col_segment])
                if col_segment is not None and len(row) > col_segment
                else ""
            )

            if not raw_dist or not raw_owner:
                continue

            if self._is_skipped_distributor(raw_dist):
                skipped_bajaj += 1
                continue

            dist_name = normalize_company_name(raw_dist)
            owner_name = " ".join(raw_owner.split())
            seg_name = self._normalize_segment(raw_seg)

            dedupe_key = (owner_name.casefold(), dist_name.casefold())
            if dedupe_key in seen_keys:
                continue
            seen_keys.add(dedupe_key)

            parsed.append(
                {
                    "distributor": dist_name,
                    "owner": owner_name,
                    "segment": seg_name,
                    "maps_to_admin": "true" if self._is_admin_owner(owner_name) else "false",
                }
            )

        if not parsed:
            raise ValidationAppError("No valid distributor mapping records were found in the uploaded file")

        logger.info(
            "Persona Excel parsed | rows={} | skipped_bajaj={}",
            len(parsed),
            skipped_bajaj,
        )
        return parsed

    def preview_import(self, file_path: Union[str, Path]) -> Dict[str, Any]:
        """Preview parsed Excel records before applying DB changes."""
        records = self.parse_excel(file_path)
        master = self._load_master_distributors()

        matched = 0
        will_create: List[str] = []
        for r in records:
            hit = self._find_master_distributor(r["distributor"], master)
            if hit:
                matched += 1
                r["matched_master"] = hit.company or hit.name
                r["will_create"] = "false"
            else:
                will_create.append(r["distributor"])
                r["matched_master"] = None
                r["will_create"] = "true"

        sales_owners = sorted(
            {r["owner"] for r in records if r.get("maps_to_admin") != "true"}
        )
        admin_rows = sum(1 for r in records if r.get("maps_to_admin") == "true")
        distinct_segments = sorted({r["segment"] for r in records if r["segment"]})

        return {
            "total_rows": len(records),
            "unique_owners": len(sales_owners),
            "new_users_count": len(sales_owners),
            "new_users": sales_owners,
            "admin_mapped_rows": admin_rows,
            "matched_distributors": matched,
            "distributors_to_create": sorted(set(will_create)),
            # Keep key for older UI — these are created, not skipped
            "unmatched_distributors": [],
            "distinct_segments": distinct_segments,
            "preview_samples": records[:50],
            "default_password": TEMP_DEFAULT_PASSWORD,
            "note": (
                "Bajaj rows are skipped. FYIP DC maps to Admin (Deb). "
                "Any Second Party not already in Master is created automatically."
            ),
        }

    @classmethod
    def _normalize_segment(cls, seg_str: str) -> str:
        """Map segment string to one of the 5 core business segments if matched."""
        if not seg_str:
            return "General"
        val = seg_str.lower()
        if "paper" in val:
            return "Paper"
        if "carpet" in val:
            return "Carpet"
        if "construct" in val:
            return "Construction"
        if "rubber" in val:
            return "Rubber"
        if "glove" in val:
            return "Gloves"
        return seg_str.title()

    def execute_import(
        self,
        file_path: Union[str, Path],
        *,
        replace_existing: bool = True,
        actor: str = "system",
        assigned_by: Optional[int] = None,
    ) -> Dict[str, Any]:
        """
        Execute bulk import from persona master Excel.

        - Soft-deletes all prior Sales Owners (keeps Super Admin + Admin)
        - Skips Bajaj only
        - Maps FYIP DC rows onto Admin (Deb) login
        - Creates any missing distributors in Master, then links them
        - New sales users: username = first name lowercase, password = Sales@123
        """
        records = self.parse_excel(file_path)
        master = self._load_master_distributors()
        admin_user = self._resolve_admin_user()

        if replace_existing:
            sales_users = self.db.scalars(
                select(User).where(
                    User.role == UserRole.USER.value,
                    User.is_deleted.is_(False),
                )
            ).all()
            sales_ids = [u.id for u in sales_users]
            for u in sales_users:
                u.is_deleted = True
                u.is_active = False

            # Clear mappings for removed sales owners + admin (re-applied from Excel)
            clear_ids = list(sales_ids)
            if admin_user:
                clear_ids.append(admin_user.id)
            if clear_ids:
                self.db.execute(delete(UserDistributor).where(UserDistributor.user_id.in_(clear_ids)))
                self.db.execute(delete(UserSegment).where(UserSegment.user_id.in_(clear_ids)))
            self.db.flush()

        existing_usernames = {
            (u.username or "").strip().lower()
            for u in self.db.scalars(select(User).where(User.is_deleted.is_(False))).all()
        }

        created_users: List[str] = []
        user_map_by_owner: dict[str, User] = {}
        pw_hash = _hash_pw(TEMP_DEFAULT_PASSWORD)

        sales_owner_names = sorted(
            {r["owner"] for r in records if r.get("maps_to_admin") != "true"}
        )
        for owner_name in sales_owner_names:
            key = owner_name.casefold()
            # Prefer reactivating soft-deleted user with same full name
            existing = self.db.scalar(
                select(User).where(
                    User.full_name.ilike(owner_name),
                    User.role == UserRole.USER.value,
                )
            )
            if existing and existing.is_deleted:
                uname = self._generate_username(owner_name, existing_usernames)
                existing.is_deleted = False
                existing.is_active = True
                existing.username = uname
                existing.email = f"{uname}@apcotex.com"
                existing.full_name = owner_name
                existing.title = "Sales Owner"
                existing.password_hash = pw_hash
                self.db.flush()
                user_map_by_owner[key] = existing
                created_users.append(uname)
            elif existing and not existing.is_deleted:
                # Fresh after replace this shouldn't happen; still reset credentials
                existing.password_hash = pw_hash
                existing.is_active = True
                user_map_by_owner[key] = existing
            else:
                uname = self._generate_username(owner_name, existing_usernames)
                new_user = User(
                    username=uname,
                    email=f"{uname}@apcotex.com",
                    full_name=owner_name,
                    title="Sales Owner",
                    password_hash=pw_hash,
                    role=UserRole.USER.value,
                    is_active=True,
                    is_deleted=False,
                    outlook_sync_permission="own",
                )
                self.db.add(new_user)
                self.db.flush()
                user_map_by_owner[key] = new_user
                created_users.append(uname)

        mapped_distributors_count = 0
        created_distributors: List[str] = []
        assigned_ud_pairs: set[tuple[int, int]] = set()
        admin_mapped = 0

        for rec in records:
            dist_name = rec["distributor"]
            owner_name = rec["owner"]
            maps_admin = rec.get("maps_to_admin") == "true"

            if maps_admin:
                user = admin_user
                if not user:
                    logger.warning("FYIP DC row skipped — no Admin user found | distributor={}", dist_name)
                    continue
                admin_mapped += 1
            else:
                user = user_map_by_owner.get(owner_name.casefold())
            if not user:
                continue

            dist, was_created = self._resolve_or_create_distributor(
                dist_name,
                owner_name=owner_name if not maps_admin else (admin_user.full_name if admin_user else "Admin"),
                master=master,
            )
            if was_created:
                created_distributors.append(dist.company or dist.name or dist_name)

            ud_pair = (user.id, dist.id)
            if ud_pair not in assigned_ud_pairs:
                assigned_ud_pairs.add(ud_pair)
                existing_ud = self.db.scalar(
                    select(UserDistributor).where(
                        UserDistributor.user_id == user.id,
                        UserDistributor.distributor_id == dist.id,
                    )
                )
                if not existing_ud:
                    self.db.add(
                        UserDistributor(
                            user_id=user.id,
                            distributor_id=dist.id,
                            assigned_by=assigned_by,
                        )
                    )
                    mapped_distributors_count += 1

        # Segment assign pass per sales owner (Admin already has all segments)
        owner_segments: dict[int, set[str]] = {}
        for rec in records:
            if rec.get("maps_to_admin") == "true":
                continue
            user = user_map_by_owner.get(rec["owner"].casefold())
            if not user:
                continue
            seg = self._normalize_segment(rec["segment"])
            if not seg or seg == "General":
                continue
            owner_segments.setdefault(user.id, set()).add(seg)
        for uid, segs in owner_segments.items():
            self.user_segments.assign_segments_to_user(
                uid, sorted(segs), assigned_by=assigned_by
            )

        self.db.flush()

        self.audit.log(
            AuditTrailCreate(
                user_name=actor,
                action=AuditAction.CREATED,
                details=(
                    f"Bulk imported persona mappings from Excel | records={len(records)} | "
                    f"created_users={len(created_users)} | mapped_distributors={mapped_distributors_count} | "
                    f"created_distributors={len(set(created_distributors))} | admin_mapped={admin_mapped}"
                ),
                entity_type="user",
                entity_id="bulk_persona_import",
            )
        )

        logger.info(
            "Bulk persona import complete | records={} | created_users={} | mapped={} | "
            "created_distributors={} | admin_rows={}",
            len(records),
            len(created_users),
            mapped_distributors_count,
            len(set(created_distributors)),
            admin_mapped,
        )

        return {
            "success": True,
            "total_records": len(records),
            "created_users_count": len(created_users),
            "created_users": created_users,
            "mapped_distributors_count": mapped_distributors_count,
            "created_distributors_count": len(set(created_distributors)),
            "created_distributors": sorted(set(created_distributors)),
            "admin_mapped_rows": admin_mapped,
            "unmatched_distributors": [],
            "default_password": TEMP_DEFAULT_PASSWORD,
            "message": (
                f"Processed {len(records)} mappings · {len(created_users)} sales logins "
                f"(password {TEMP_DEFAULT_PASSWORD}) · {mapped_distributors_count} distributor links"
                + (
                    f" · {len(set(created_distributors))} new distributors added to Master"
                    if created_distributors
                    else ""
                )
                + (f" · {admin_mapped} FYIP DC → Admin" if admin_mapped else "")
            ),
        }

    @staticmethod
    def generate_sample_excel() -> Path:
        """Generate sample Excel file for distributor-employee mapping download."""
        from app.core.config import settings

        out_dir = Path(settings.download_dir) / "templates"
        ensure_dir(out_dir)
        path = out_dir / "bulk_persona_import_template.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Employee Mapping"

        headers = ["Second Party", "Owner", "Segment"]
        ws.append(headers)

        sample_rows = [
            ["Chaudhury", "Harsh", "Rubber"],
            ["Reda Industries", "Anup Pandey", "Construction"],
            ["Altek", "Sachin Kasar", "Rubber"],
            ["FYIP DC Partner", "FYIP DC", "Rubber"],
        ]
        for row in sample_rows:
            ws.append(row)

        for col in range(1, 4):
            ws.column_dimensions[ws.cell(1, col).column_letter].width = 24

        wb.save(path)
        return path
