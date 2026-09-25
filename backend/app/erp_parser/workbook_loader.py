"""Load .xlsx, .xlsm, and .xls into one workbook surface.

.xlsx/.xlsm use openpyxl. .xls uses xlrd. Callers keep using sheet names,
cell values, merged ranges, and datetimes. The file on disk is not converted.
"""

from __future__ import annotations

import time
from pathlib import Path
from typing import Any, Iterator, List, Optional, Sequence, Union

import xlrd
from openpyxl import load_workbook
from xlrd import XL_CELL_BLANK, XL_CELL_BOOLEAN, XL_CELL_DATE, XL_CELL_EMPTY, XL_CELL_ERROR, XL_CELL_NUMBER
from xlrd.sheet import Sheet

from app.core.logging import get_logger
from app.exceptions import ExcelProcessingError

logger = get_logger(__name__)

ALLOWED_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}
_OPENPYXL_EXTENSIONS = {".xlsx", ".xlsm"}

# One success line per file while a parse re-opens the same workbook.
_LOGGED_AT: dict[str, float] = {}
_LOG_WINDOW_SECONDS = 8.0


class _Cell:
    def __init__(self, value: Any) -> None:
        self.value = value


class _Merge:
    """Inclusive 1-based bounds, matching openpyxl merged ranges."""

    def __init__(self, rlo: int, rhi: int, clo: int, chi: int) -> None:
        self.min_row = rlo + 1
        self.max_row = rhi
        self.min_col = clo + 1
        self.max_col = chi


class _Merges:
    def __init__(self, ranges: Sequence[_Merge]) -> None:
        self.ranges = list(ranges)


class _ColDim:
    def __init__(self, hidden: bool) -> None:
        self.hidden = hidden


class _ColDims:
    def __init__(self, sheet: Sheet) -> None:
        self._sheet = sheet

    def get(self, letter: str) -> _ColDim:
        index = _column_index(letter) - 1
        info = getattr(self._sheet, "colinfo_map", {}).get(index)
        hidden = bool(getattr(info, "hidden", False)) if info is not None else False
        return _ColDim(hidden)


class XlsSheet:
    """xlrd worksheet with the cell, merge, and visibility surface the pipeline reads."""

    def __init__(self, book: xlrd.Book, sheet: Sheet) -> None:
        self._book = book
        self._sheet = sheet
        self.title = sheet.name
        self.merged_cells = _Merges(_merge_ranges(sheet))
        self.column_dimensions = _ColDims(sheet)
        visibility = int(getattr(sheet, "visibility", 0) or 0)
        if visibility == 1:
            self.sheet_state = "hidden"
        elif visibility >= 2:
            self.sheet_state = "veryHidden"
        else:
            self.sheet_state = "visible"

    def cell(self, row: int, column: int) -> _Cell:
        return _Cell(self._value(row - 1, column - 1))

    def iter_rows(
        self,
        min_row: int = 1,
        max_row: Optional[int] = None,
        min_col: int = 1,
        max_col: Optional[int] = None,
        values_only: bool = False,
    ) -> Iterator[tuple]:
        last_row = self._sheet.nrows if max_row is None else min(max_row, self._sheet.nrows)
        width = max_col if max_col is not None else self._sheet.ncols
        start_row = max(1, min_row)
        start_col = max(1, min_col)
        if last_row < start_row:
            return
        for row_idx in range(start_row, last_row + 1):
            values = [self._value(row_idx - 1, col_idx - 1) for col_idx in range(start_col, width + 1)]
            if values_only:
                yield tuple(values)
            else:
                yield tuple(_Cell(value) for value in values)

    def _value(self, row: int, col: int) -> Any:
        if row < 0 or col < 0 or row >= self._sheet.nrows or col >= self._sheet.ncols:
            return None
        kind = self._sheet.cell_type(row, col)
        if kind in (XL_CELL_EMPTY, XL_CELL_BLANK, XL_CELL_ERROR):
            return None
        raw = self._sheet.cell_value(row, col)
        if kind == XL_CELL_DATE:
            try:
                return xlrd.xldate_as_datetime(raw, self._book.datemode)
            except Exception:  # noqa: BLE001
                return raw
        if kind == XL_CELL_NUMBER:
            if isinstance(raw, float) and raw.is_integer():
                return int(raw)
            return raw
        if kind == XL_CELL_BOOLEAN:
            return bool(raw)
        if isinstance(raw, str):
            return raw
        return raw


class XlsWorkbook:
    """xlrd book exposed with sheetnames, item access, and close()."""

    def __init__(self, book: xlrd.Book) -> None:
        self._book = book
        self._sheets = [XlsSheet(book, sheet) for sheet in book.sheets()]
        self.sheetnames = [sheet.title for sheet in self._sheets]
        self._by_name = {sheet.title: sheet for sheet in self._sheets}

    def __getitem__(self, name: str) -> XlsSheet:
        return self._by_name[name]

    def close(self) -> None:
        release = getattr(self._book, "release_resources", None)
        if callable(release):
            release()


def load_workbook_any(path: Union[str, Path]):
    """Open a workbook. .xls uses xlrd. .xlsx and .xlsm use openpyxl."""
    file_path = Path(path)
    if not file_path.exists():
        raise ExcelProcessingError(f"Excel file not found: {file_path}")
    suffix = file_path.suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ExcelProcessingError(
            f"Unsupported workbook type '{suffix}'. Only .xlsx, .xlsm, and .xls are accepted."
        )
    try:
        if suffix == ".xls":
            workbook = XlsWorkbook(_open_xls(file_path))
            engine = "xlrd"
        elif suffix in _OPENPYXL_EXTENSIONS:
            workbook = load_workbook(file_path, data_only=True, read_only=False)
            engine = "openpyxl"
        else:
            raise ExcelProcessingError(
                f"Unsupported workbook type '{suffix}'. Only .xlsx, .xlsm, and .xls are accepted."
            )
    except ExcelProcessingError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExcelProcessingError(f"Unable to open workbook: {exc}") from exc

    _log_loaded(file_path, engine, len(list(workbook.sheetnames)))
    return workbook


def _open_xls(path: Path) -> xlrd.Book:
    try:
        return xlrd.open_workbook(str(path), formatting_info=True)
    except Exception:
        return xlrd.open_workbook(str(path), formatting_info=False)


def _log_loaded(path: Path, engine: str, sheet_count: int) -> None:
    key = str(path.resolve())
    now = time.monotonic()
    previous = _LOGGED_AT.get(key)
    _LOGGED_AT[key] = now
    if previous is not None and now - previous < _LOG_WINDOW_SECONDS:
        return
    logger.info(
        "Workbook Loader\n\nFile : {}\n\nFormat : {}\n\nEngine : {}\n\nSheets : {}\n\nStatus : Loaded Successfully",
        path.name,
        path.suffix.lstrip(".").upper(),
        engine,
        sheet_count,
    )


def _merge_ranges(sheet: Sheet) -> List[_Merge]:
    ranges = getattr(sheet, "merged_cells", None) or []
    return [_Merge(rlo, rhi, clo, chi) for rlo, rhi, clo, chi in ranges]


def _column_index(letter: str) -> int:
    index = 0
    for char in (letter or "").strip().upper():
        if not ("A" <= char <= "Z"):
            break
        index = index * 26 + (ord(char) - 64)
    return index or 1
