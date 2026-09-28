"""Word tables to a normalized grid. Paragraphs are kept only when they carry period metadata."""

from __future__ import annotations

import re
from pathlib import Path

from app.erp_parser.documents.grid import GridSheet, NormalizedGrid
from app.exceptions import ExcelProcessingError

_META_RE = re.compile(
    r"\b(?:FY\s*\d{4}|Q\s*[1-4]|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC|JAN|FEB|MAR)\b",
    re.IGNORECASE,
)


def docx_to_grid(path: Path) -> NormalizedGrid:
    try:
        from docx import Document
    except ImportError as exc:
        raise ExcelProcessingError("python-docx is not installed") from exc

    document = Document(str(path))
    metadata = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text and _META_RE.search(paragraph.text)
    ][:5]
    sheets: list[GridSheet] = []
    for index, table in enumerate(document.tables, start=1):
        rows = []
        if index == 1:
            rows.extend([[line] for line in metadata])
        for raw in table.rows:
            cells = [cell.text.strip() or None for cell in raw.cells]
            if any(cells):
                rows.append(cells)
        if rows:
            sheets.append(GridSheet(name=f"Table {index}", rows=rows))
    if not sheets:
        raise ExcelProcessingError(f"No tables found in '{path.name}'")
    return NormalizedGrid(document_type="DOCX", sheets=sheets)
