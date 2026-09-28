"""Write a normalized grid to a temporary workbook the orchestrator already understands."""

from __future__ import annotations

import tempfile
from pathlib import Path

from openpyxl import Workbook

from app.erp_parser.documents.grid import NormalizedGrid
from app.exceptions import ExcelProcessingError


def materialize_grid(grid: NormalizedGrid, *, source_name: str) -> Path:
    """Persist the grid as xlsx. The caller deletes the file after parsing."""
    populated = [sheet for sheet in grid.sheets if sheet.rows]
    if not populated:
        raise ExcelProcessingError(f"No tables found in '{source_name}'")
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    used: set[str] = set()
    for index, sheet in enumerate(populated, start=1):
        title = _sheet_title(sheet.name, index, used)
        worksheet = workbook.create_sheet(title)
        for row in sheet.rows:
            worksheet.append([_cell(value) for value in row])
        for min_row, max_row, min_col, max_col in sheet.merged:
            if max_row >= min_row and max_col >= min_col:
                worksheet.merge_cells(
                    start_row=min_row,
                    end_row=max_row,
                    start_column=min_col,
                    end_column=max_col,
                )
    handle = tempfile.NamedTemporaryFile(
        prefix=f"{Path(source_name).stem[:40]}-",
        suffix=".xlsx",
        delete=False,
    )
    handle.close()
    target = Path(handle.name)
    workbook.save(target)
    workbook.close()
    return target


def _sheet_title(name: str, index: int, used: set[str]) -> str:
    raw = "".join(ch for ch in (name or f"Page {index}") if ch not in "[]:*?/\\")
    title = (raw or f"Page {index}")[:31]
    candidate = title
    suffix = 2
    while candidate in used:
        extra = f" {suffix}"
        candidate = f"{title[: 31 - len(extra)]}{extra}"
        suffix += 1
    used.add(candidate)
    return candidate


def _cell(value: object) -> object:
    if value is None:
        return None
    text = str(value).strip()
    return text or None
