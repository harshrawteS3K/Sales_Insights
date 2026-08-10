"""Helpers to build official APCOTEX quarterly sales workbooks for tests."""

from pathlib import Path
from typing import Optional, Sequence

from openpyxl import Workbook


def build_official_workbook(
    path: Path,
    *,
    distributor: str,
    reporting_month: str,
    rows: Sequence[tuple],
    include_stock: bool = False,
    legacy_period_column: bool = False,
    company: Optional[str] = None,
    address: str = "Addr",
    phone: str = "999",
    use_legacy_month_labels: bool = False,
) -> Path:
    """
    Build the finalized quarterly official template.

    Metadata: Name of Person, Company Name, Reporting Quarter
    Sales: Sr. No., Customer Name, Segment, Product, Sales Quantity

    ``include_stock`` / legacy labels kept for backward-compat tests only.
    """
    wb = Workbook()
    ws = wb.active
    company_value = company if company is not None else distributor

    if use_legacy_month_labels:
        ws["A1"] = "Name of Distributor"
        ws["B1"] = distributor
        ws["A2"] = "Company Name"
        ws["B2"] = company_value
        ws["A3"] = "Address"
        ws["B3"] = address
        ws["A4"] = "Phone No"
        ws["B4"] = phone
        ws["A5"] = "Reporting Month"
        ws["B5"] = reporting_month
        header_row = 9
        headers = ["Sr. No.", "Name of Customer", "Segment", "Product", "Quantity"]
    else:
        ws["A1"] = "Name of Person"
        ws["B1"] = distributor
        ws["A2"] = "Company Name"
        ws["B2"] = company_value
        ws["A3"] = "Reporting Quarter"
        ws["B3"] = reporting_month
        header_row = 6
        headers = [
            "Sr. No.",
            "Customer Name",
            "Segment",
            "Product",
            "Sales Quantity",
        ]

    if include_stock:
        headers.extend(["Opening Stock", "Closing Stock"])
    if legacy_period_column:
        headers.append("Period")

    for i, h in enumerate(headers, 1):
        ws.cell(header_row, i, h)

    # Only write as many cells as headers (ignore trailing stock when include_stock=False)
    width = len(headers)
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(list(row)[:width], 1):
            ws.cell(header_row + r, c, v)

    wb.save(path)
    return path
