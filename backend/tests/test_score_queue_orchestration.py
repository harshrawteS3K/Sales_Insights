"""Targeted tests: score-queue orchestration + Avik / A.B. Brothers layouts."""

from __future__ import annotations

import threading
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook

from app.erp_parser import ERPParserService
from app.erp_parser.header_dictionary import reload_header_dictionary
from app.enums import EmailProcessStatus
from app.services import erp_score_queue as score_q


def _save(wb: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


@pytest.fixture(autouse=True)
def _reload_headers():
    reload_header_dictionary()
    yield


def _avik_workbook(path: Path, *, month: str = "April") -> Path:
    """Two-sheet Avik-style: cover sheet + Party/Item/Billed Quantity/Month."""
    wb = Workbook()
    cover = wb.active
    cover.title = "Sales 26-27"
    cover.append(["Summary only"])
    cover.append(["Not the transaction table"])

    data = wb.create_sheet("Sales 26-27 (2)")
    data.append(["Party Name", "Item Name", "Billed Quantity", "Month"])
    data.append(["POLYMEK PRODUCTS", "N 745", 210, month])
    data.append(["SHRIRAM RUBBER PRODUCTS PVT.LTD", "N 745", 525, month])
    data.append(["TREEMURTI RUBBERCRAFT", "N 745", 105, month])
    return _save(wb, path)


def _ab_brothers_workbook(path: Path) -> Path:
    """Sheet1 noise + Data sheet with Party/Item/Unit/Total."""
    wb = Workbook()
    cover = wb.active
    cover.title = "Sheet1"
    cover.append(["Row Labels"])
    cover.append(["Some summary"])
    cover.append(["Grand Total", 99999])

    data = wb.create_sheet("Data")
    data.append(["Name of the Party", "Name of the Item", "Unit", "Total"])
    data.append(["ACCURATE RUB TECH", "APCOFLEX NVC573E", "Kgs", 175])
    data.append(["ACCURATE RUB TECH", "APCOFLEX N745", "Kgs", 875])
    data.append(["AMBICA RUBBER (AHEMDABAD)", "APCOFLEX N745", "Kgs", 1050])
    return _save(wb, path)


def test_avik_style_two_sheet_retains_confidence(tmp_path: Path):
    """A) Avik-style two-sheet → valid sheet discovered, rows parsed, confidence retained."""
    path = _avik_workbook(tmp_path / "Avik_April.xlsx", month="April")
    result = ERPParserService().parse_workbook(path, allow_llm_fallback=False)
    assert result.imported_rows == 3
    assert float(result.overall_confidence) >= 70
    assert result.rows[0]["customer_name"] == "POLYMEK PRODUCTS"
    assert result.rows[0]["product"] == "N 745"
    assert float(result.rows[0]["sales_quantity"]) == 210
    assert "Sales 26-27 (2)" in str(result.sheet_name)


def test_ab_brothers_data_sheet_discovered(tmp_path: Path):
    """B) A.B. Brothers two-sheet → Data sheet discovered, rows parsed."""
    path = _ab_brothers_workbook(tmp_path / "WEST_AB_BROTHER.xlsx")
    result = ERPParserService().parse_workbook(path, allow_llm_fallback=False)
    assert result.imported_rows == 3
    assert float(result.overall_confidence) >= 70
    customers = {r["customer_name"] for r in result.rows}
    assert "ACCURATE RUB TECH" in customers
    assert "AMBICA RUBBER (AHEMDABAD)" in customers
    qtys = {float(r["sales_quantity"]) for r in result.rows}
    assert 175.0 in qtys
    assert 875.0 in qtys
    assert 1050.0 in qtys
    assert str(result.sheet_name) == "Data"


def test_three_avik_attachments_independent(tmp_path: Path):
    """C) April/May/June processed independently — one failure cannot wipe others."""
    april = _avik_workbook(tmp_path / "April.xlsx", month="April")
    may = _avik_workbook(tmp_path / "May.xlsx", month="May")
    june = _avik_workbook(tmp_path / "June.xlsx", month="June")

    progress = score_q._empty_progress()
    lock = threading.Lock()

    class FakeAtt:
        def __init__(self, path: Path, fail: bool = False):
            self.file_name = path.name
            self.file_path = str(path) if not fail else str(path) + ".missing"
            self.is_excel = True
            self.is_deleted = False

    class FakeEmail:
        id = 1
        process_status = EmailProcessStatus.SCORING.value
        confidence_score = 0
        parsed_distributor = "Distributor"
        detected_quarter = "FY 2025-26 • Q1"
        subject = "Sales Q1"
        parsed_segment = "West"
        attachments = [FakeAtt(april), FakeAtt(may), FakeAtt(june, fail=True)]
        mapping_source = None

    fake_email = FakeEmail()
    db = MagicMock()

    def _scalar(stmt):  # noqa: ARG001
        return fake_email

    db.scalar.side_effect = _scalar

    with patch("app.services.erp_score_queue.detect_reporting_quarter", create=True):
        with patch.object(score_q, "_remember_quarter", return_value="subject_or_existing"):
            outcome = score_q._run_score(
                db,
                1,
                allow_llm=False,
                progress=progress,
                progress_lock=lock,
            )

    assert outcome["process_status"] == EmailProcessStatus.PARSED.value
    assert int(outcome["confidence"]) >= 70
    successes = [a for a in progress["attachments"] if a["status"] == "success"]
    failures = [a for a in progress["attachments"] if a["status"] != "success"]
    assert len(successes) == 2
    assert len(failures) == 1
    assert int(outcome["confidence"]) == int(round(min(a["confidence"] for a in successes)))


def test_parser_100_plus_timeout_retains_confidence():
    """D) Parser confidence=100 + enrichment timeout → must NOT become 0."""
    progress = score_q._empty_progress()
    progress["attachments"] = [
        {
            "file": "Avik_June.xlsx",
            "status": "success",
            "confidence": 100.0,
            "rows": 139,
            "mapping_source": "python",
        }
    ]
    progress["warnings"] = ["Scoring enrichment timed out"]
    progress["quarter_status"] = "timeout"

    class FakeEmail:
        id = 166
        process_status = EmailProcessStatus.SCORING.value
        confidence_score = 100
        mapping_source = "python"

    outcome = score_q._outcome_from_progress(progress)
    assert int(outcome["confidence"]) == 100
    assert outcome["process_status"] == EmailProcessStatus.PARSED.value

    protected = score_q._protect_successful_parse(
        FakeEmail(),
        {
            "process_status": EmailProcessStatus.FAILED.value,
            "confidence": 0,
            "error_message": "Scoring timeout",
            "mapping_source": None,
        },
    )
    assert int(protected["confidence"]) == 100
    assert protected["process_status"] == EmailProcessStatus.PARSED.value
    assert "retained" in (protected.get("error_message") or "").lower() or "timeout" in (
        protected.get("error_message") or ""
    ).lower()


def test_low_confidence_attempts_llm_fallback(tmp_path: Path):
    """E) Low-confidence workbook → LLM fallback is attempted."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Who Bought It", "Thing Sold SKU", "Bags Moved"])
    ws.append(["Omega Mills", "Latex-100", 55])
    path = _save(wb, tmp_path / "low.xlsx")

    mock = MagicMock()
    mock.resolve_headers.return_value = {
        "customer_column": "Who Bought It",
        "product_column": "Thing Sold SKU",
        "quantity_column": "Bags Moved",
        "confidence": 96,
        "layout_type": "standard_header",
        "parser_hint": "header",
        "is_sales_table": True,
    }
    mock.last_token_count = 1
    mock.last_prompt_tokens = 1
    mock.last_completion_tokens = 1
    mock.last_model_id = "test"

    with patch(
        "app.erp_parser.orchestrator.service.LLMHeaderResolver",
        return_value=mock,
    ):
        result = ERPParserService().parse_workbook(path, allow_llm_fallback=True)

    mock.resolve_headers.assert_called()
    assert result.imported_rows == 1
    assert result.mapping_source == "llm"


def test_llm_empty_quarter_preserves_deterministic_parse(tmp_path: Path):
    """F) LLM empty quarter response → warning only; deterministic parse preserved."""
    path = _avik_workbook(tmp_path / "Avik_May.xlsx", month="May")
    result = ERPParserService().parse_workbook(path, allow_llm_fallback=False)
    assert result.imported_rows == 3
    conf = float(result.overall_confidence)
    assert conf >= 70

    progress = score_q._empty_progress()
    progress["attachments"] = [
        {
            "file": "Avik_May.xlsx",
            "status": "success",
            "confidence": conf,
            "rows": 3,
            "mapping_source": "python",
        }
    ]
    progress["quarter_status"] = "warning"
    progress["warnings"] = ["Quarter detection: warning"]

    outcome = score_q._outcome_from_progress(progress)
    assert int(outcome["confidence"]) == int(round(conf))
    assert outcome["process_status"] == EmailProcessStatus.PARSED.value
    assert "Quarter" in (outcome.get("error_message") or "")


def test_one_failed_attachment_keeps_others_successful(tmp_path: Path):
    """G) Multiple attachments where one fails → successful ones remain successful."""
    good = _avik_workbook(tmp_path / "good.xlsx", month="April")
    progress = score_q._empty_progress()
    lock = threading.Lock()

    class FakeAtt:
        def __init__(self, path: Path | None, name: str):
            self.file_name = name
            self.file_path = str(path) if path else None
            self.is_excel = True
            self.is_deleted = False

    class FakeEmail:
        id = 2
        process_status = EmailProcessStatus.SCORING.value
        confidence_score = 0
        parsed_distributor = ""
        detected_quarter = "FY 2025-26 • Q1"
        subject = "Sales"
        parsed_segment = ""
        attachments = [
            FakeAtt(good, "good.xlsx"),
            FakeAtt(None, "bad.xlsx"),
        ]
        mapping_source = None

    db = MagicMock()
    db.scalar.return_value = FakeEmail()

    with patch.object(score_q, "_remember_quarter", return_value="subject_or_existing"):
        outcome = score_q._run_score(
            db,
            2,
            allow_llm=False,
            progress=progress,
            progress_lock=lock,
        )

    assert outcome["process_status"] == EmailProcessStatus.PARSED.value
    assert int(outcome["confidence"]) >= 70
    assert any(a["status"] == "success" for a in progress["attachments"])
    assert any(a["status"] != "success" for a in progress["attachments"])


def test_persist_timeout_preserves_checkpoint(tmp_path: Path):
    """Timeout after checkpoint must retain parser confidence, not write 0."""
    progress = score_q._empty_progress()
    progress["attachments"] = [
        {
            "file": "Avik_June.xlsx",
            "status": "success",
            "confidence": 100.0,
            "rows": 139,
            "mapping_source": "python",
        }
    ]
    captured = {}

    def fake_persist(email_id, *, process_status, confidence, error_message, mapping_source=None):
        captured.update(
            {
                "email_id": email_id,
                "process_status": process_status,
                "confidence": confidence,
                "error_message": error_message,
                "mapping_source": mapping_source,
            }
        )

    with patch.object(score_q, "_persist_terminal", side_effect=fake_persist):
        score_q._persist_terminal_preserving_success(
            166,
            fallback_status=EmailProcessStatus.FAILED.value,
            fallback_confidence=0,
            error_message="Scoring timeout",
            progress=progress,
        )

    assert captured["confidence"] == 100
    assert captured["process_status"] == EmailProcessStatus.PARSED.value
    assert captured["confidence"] != 0
