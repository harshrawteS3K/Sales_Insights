"""Prior-quarter customer selection for distributor-specific templates."""

from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database.session import SessionLocal
from app.repositories.distributor_customer_mapping_repository import (
    DistributorCustomerMappingRepository,
)
from app.repositories.distributor_repository import DistributorRepository
from app.schemas.distributor import TemplateGenerateRequest
from app.services.master_data_service import MasterDataService
from app.services.product_master_service import ProductMasterService
from app.services.report_service import ReportService
from app.services.template_generation_service import TemplateGenerationService
from app.utils.period_calendar import (
    is_quarter_strictly_before,
    previous_financial_quarter,
    quarter_sort_key,
)
from tests.workbook_helpers import build_official_workbook


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
        session.rollback()
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()


def _uid(prefix: str) -> str:
    return f"{prefix} {uuid4().hex[:8]}"


def _seed_products(db: Session) -> None:
    try:
        ProductMasterService(db).replace_from_rows(
            [("Paper", "P1"), ("Paper", "P2")], actor="test"
        )
    except Exception:
        from app.models.product_master import ProductMaster

        for industry, code in [("Paper", "P1"), ("Paper", "P2")]:
            db.add(
                ProductMaster(
                    industry_type=industry,
                    product_code=code,
                    is_active=True,
                )
            )
        db.flush()


def test_previous_financial_quarter_helpers():
    assert previous_financial_quarter("Q1 2027") == "Q4 2026"
    assert previous_financial_quarter("Q2 2026") == "Q1 2026"
    assert previous_financial_quarter("Q3 2026") == "Q2 2026"
    assert previous_financial_quarter("Q4 2026") == "Q3 2026"
    assert quarter_sort_key("Q1 2026") < quarter_sort_key("Q2 2026")
    assert is_quarter_strictly_before("Q1 2026", "Q3 2026")
    assert is_quarter_strictly_before("Q4 2026", "Q1 2027")
    assert not is_quarter_strictly_before("Q3 2026", "Q3 2026")
    assert not is_quarter_strictly_before("Q3 2026", "Q2 2026")


def test_template_includes_all_historical_customers(db: Session, tmp_path: Path):
    """All imported quarters contribute customers; latest quarter ordered first."""
    _seed_products(db)
    company = _uid("Prior Co")

    q1 = build_official_workbook(
        tmp_path / "q1.xlsx",
        distributor="Rep Prior",
        company=company,
        reporting_month="Q1 2095",
        rows=[(1, "AKS RUGS", "Paper", "P1", 10)],
    )
    q2 = build_official_workbook(
        tmp_path / "q2.xlsx",
        distributor="Rep Prior",
        company=company,
        reporting_month="Q2 2095",
        rows=[(1, "XYZ POLYMERS", "Paper", "P1", 5)],
    )
    ReportService(db).ingest_excel(q1, actor="tester")
    ReportService(db).ingest_excel(q2, actor="tester")

    dist = DistributorRepository(db).get_by_company(company)
    assert dist is not None

    names = DistributorCustomerMappingRepository(db).list_names_for_template(dist.id)
    assert any(n.casefold() == "aks rugs" for n in names)
    assert any(n.casefold() == "xyz polymers" for n in names)
    # Latest quarter (Q2) first
    assert names[0].casefold() == "xyz polymers"

    out = tmp_path / "template.xlsx"
    result = TemplateGenerationService(db).generate(
        actor="tester",
        mode="distributor",
        distributor_id=dist.id,
        reporting_quarter=None,
        output_path=out,
        update_canonical=False,
    )
    assert result.fallback_generic is False
    assert result.customers_count == 2
    wb = load_workbook(out)
    assert "CustomerList" in wb.sheetnames
    cust_sheet = wb["CustomerList"]
    values = {
        str(cust_sheet.cell(row=r, column=1).value or "").strip().casefold()
        for r in range(1, cust_sheet.max_row + 1)
    }
    assert "aks rugs" in values
    assert "xyz polymers" in values


def test_no_history_falls_back_to_generic(db: Session, tmp_path: Path):
    _seed_products(db)
    from app.schemas.distributor import DistributorCreate
    from app.services.distributor_service import DistributorService

    dist = DistributorService(db).create_distributor(
        DistributorCreate(
            name=_uid("Bare"),
            company=_uid("Bare Co"),
            email=f"{uuid4().hex[:8]}@example.com",
        ),
        actor="admin",
    )
    out = tmp_path / "generic_fallback.xlsx"
    result = TemplateGenerationService(db).generate(
        actor="tester",
        mode="distributor",
        distributor_id=dist.id,
        reporting_quarter=None,
        output_path=out,
        update_canonical=False,
    )
    assert result.fallback_generic is True
    assert result.customers_count == 0
    assert result.warning
    assert "No customer history" in (result.warning or "")
    wb = load_workbook(out)
    assert "CustomerList" not in wb.sheetnames


def test_master_data_distributor_mode_without_reporting_quarter(
    db: Session, tmp_path: Path
):
    _seed_products(db)
    company = _uid("No RQ")
    path = build_official_workbook(
        tmp_path / "one.xlsx",
        distributor="Rep",
        company=company,
        reporting_month="Q3 2087",
        rows=[(1, "ONLY CUST", "Paper", "P1", 2)],
    )
    ReportService(db).ingest_excel(path, actor="tester")
    dist = DistributorRepository(db).get_by_company(company)
    assert dist is not None
    result = MasterDataService(db).generate_template(
        actor="admin",
        request=TemplateGenerateRequest(mode="distributor", distributor_id=dist.id),
    )
    assert result.success
    assert result.customers_count == 1
    assert result.fallback_generic is False
