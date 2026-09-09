"""LLM header resolver fallback tests (Python remains primary)."""

from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook

from app.erp_parser import ERPParserService
from app.erp_parser.llm_header_resolver import (
    LLMHeaderResolver,
    LLMHeaderResolverError,
    estimate_payload_tokens,
    parse_llm_mapping_json,
    sanitize_headers,
    sanitize_sample_rows,
)
from app.erp_parser.parser_service import LLM_FALLBACK_THRESHOLD, LLM_MIN_TRUST


def _save(wb: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def test_known_headers_python_only(tmp_path: Path):
    """Known headers → Python only; LLM must not be called."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Party Name", "Material", "Dispatch Qty"])
    ws.append(["AKS RUGS", "P1", 10])
    ws.append(["Beta Corp", "P2", 25.5])
    path = _save(wb, tmp_path / "known.xlsx")

    with patch(
        "app.erp_parser.parser_service.LLMHeaderResolver"
    ) as resolver_cls:
        result = ERPParserService().parse_workbook(path)
        resolver_cls.assert_not_called()

    assert result.mapping_source == "python"
    assert result.overall_confidence >= LLM_FALLBACK_THRESHOLD
    assert result.imported_rows == 2


def test_high_confidence_llm_not_called(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Customer Name", "Product", "Sales Quantity"])
    ws.append(["A", "X", 1])
    path = _save(wb, tmp_path / "high.xlsx")

    called = {"n": 0}

    class Spy:
        def resolve_headers(self, headers, sample_rows):
            called["n"] += 1
            raise AssertionError("LLM should not be called for high confidence")

    with patch("app.erp_parser.parser_service.LLMHeaderResolver", return_value=Spy()):
        result = ERPParserService().parse_workbook(path)

    assert called["n"] == 0
    assert result.mapping_source == "python"
    assert result.overall_confidence >= 85


def test_weird_headers_llm_called(tmp_path: Path):
    """Obscure headers → Python confidence low / incomplete → LLM called."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Who Bought It", "Thing Sold SKU", "Bags Moved"])
    ws.append(["Omega Mills", "Latex-100", 55])
    ws.append(["Acme Rugs", "NBR-745", 10])
    path = _save(wb, tmp_path / "weird.xlsx")

    llm_payload = {
        "customer_column": "Who Bought It",
        "product_column": "Thing Sold SKU",
        "quantity_column": "Bags Moved",
        "confidence": 96,
    }
    mock_resolver = MagicMock()
    mock_resolver.resolve_headers.return_value = llm_payload

    with patch(
        "app.erp_parser.parser_service.LLMHeaderResolver",
        return_value=mock_resolver,
    ):
        result = ERPParserService().parse_workbook(path)

    mock_resolver.resolve_headers.assert_called_once()
    args = mock_resolver.resolve_headers.call_args[0]
    assert len(args[0]) <= 15  # headers
    assert len(args[1]) <= 3  # sample rows
    assert result.mapping_source == "llm"
    assert result.imported_rows == 2
    assert result.rows[0]["customer_name"] == "Omega Mills"
    assert float(result.rows[0]["sales_quantity"]) == 0.055  # 55 KG → MT
    assert result.overall_confidence >= 85


def test_openai_failure_graceful_fallback(tmp_path: Path):
    """OpenAI failure → keep Python mapping; do not fail parse when Python works."""
    wb = Workbook()
    ws = wb.active
    # Fuzzy-ish headers that Python can still map but may score < 85 sometimes;
    # force LLM path by patching overall below threshold after python extract.
    ws.append(["Account Name", "FG Code", "Billed Qty"])
    ws.append(["Acme", "G7", 3])
    path = _save(wb, tmp_path / "fallback.xlsx")

    mock_resolver = MagicMock()
    mock_resolver.resolve_headers.side_effect = LLMHeaderResolverError("rate limit")

    with patch(
        "app.erp_parser.parser_service.LLMHeaderResolver",
        return_value=mock_resolver,
    ), patch(
        "app.erp_parser.parser_service.compute_erp_confidence",
        side_effect=[
            # First call: python confidence low → triggers LLM
            {
                "overall_confidence": 80.0,
                "accuracy": 80.0,
                "customer_confidence": 90.0,
                "product_confidence": 90.0,
                "quantity_confidence": 90.0,
                "band": "yellow",
                "components": {},
                "weights": {},
            },
            # Should not reach second if LLM fails — but if re-extract happens, allow
            {
                "overall_confidence": 80.0,
                "accuracy": 80.0,
                "customer_confidence": 90.0,
                "product_confidence": 90.0,
                "quantity_confidence": 90.0,
                "band": "yellow",
                "components": {},
                "weights": {},
            },
        ],
    ):
        result = ERPParserService().parse_workbook(path)

    mock_resolver.resolve_headers.assert_called_once()
    assert result.mapping_source == "python"
    assert result.imported_rows == 1
    assert result.rows[0]["customer_name"] == "Acme"


def test_json_parsing():
    raw = """```json
    {
      "customer_column": "Sold To",
      "product_column": "Commodity",
      "quantity_column": "Net Movement",
      "confidence": 96
    }
    ```"""
    parsed = parse_llm_mapping_json(raw)
    assert parsed == {
        "customer_column": "Sold To",
        "product_column": "Commodity",
        "quantity_column": "Net Movement",
        "confidence": 96,
    }

    nullish = parse_llm_mapping_json(
        '{"customer_column":null,"product_column":"P","quantity_column":"","confidence":50}'
    )
    assert nullish["customer_column"] is None
    assert nullish["quantity_column"] is None
    assert nullish["product_column"] == "P"

    with pytest.raises(LLMHeaderResolverError):
        parse_llm_mapping_json("not-json")


def test_confidence_recalculation(tmp_path: Path):
    """Blended overall = max(python, llm) when LLM succeeds."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Who Bought It", "Thing Sold SKU", "Bags Moved"])
    # Avoid short tokens like "Cust"/"Prod" that can be mis-scored as the header row
    ws.append(["Omega Mills Pvt", "Latex-100 Grade", 9])
    path = _save(wb, tmp_path / "blend.xlsx")

    mock_resolver = MagicMock()
    mock_resolver.resolve_headers.return_value = {
        "customer_column": "Who Bought It",
        "product_column": "Thing Sold SKU",
        "quantity_column": "Bags Moved",
        "confidence": 96,
    }

    with patch(
        "app.erp_parser.parser_service.LLMHeaderResolver",
        return_value=mock_resolver,
    ):
        result = ERPParserService().parse_workbook(path)

    assert result.mapping_source == "llm"
    assert result.overall_confidence >= 96 or result.overall_confidence >= LLM_MIN_TRUST
    assert result.confidence_breakdown.get("mapping_source") == "llm"
    # max(python, llm) — at least LLM confidence floor when python was weak/incomplete
    assert result.overall_confidence >= 96


def test_weak_llm_keeps_python(tmp_path: Path):
    """LLM confidence < 70 → keep parser result."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Account Name", "FG Code", "Billed Qty"])
    ws.append(["Acme", "G7", 3])
    path = _save(wb, tmp_path / "weak_llm.xlsx")

    mock_resolver = MagicMock()
    mock_resolver.resolve_headers.return_value = {
        "customer_column": "Account Name",
        "product_column": "FG Code",
        "quantity_column": "Billed Qty",
        "confidence": 40,
    }

    with patch(
        "app.erp_parser.parser_service.LLMHeaderResolver",
        return_value=mock_resolver,
    ), patch(
        "app.erp_parser.parser_service.compute_erp_confidence",
        return_value={
            "overall_confidence": 80.0,
            "accuracy": 80.0,
            "customer_confidence": 90.0,
            "product_confidence": 90.0,
            "quantity_confidence": 90.0,
            "band": "yellow",
            "components": {},
            "weights": {},
        },
    ):
        result = ERPParserService().parse_workbook(path)

    mock_resolver.resolve_headers.assert_called_once()
    assert result.mapping_source == "python"
    assert result.imported_rows == 1


def test_token_sanitization_limits():
    headers = [f"H{i}" for i in range(30)]
    clean = sanitize_headers(headers)
    assert len(clean) == 15

    rows = [[f"v{i}-{j}" for j in range(20)] for i in range(10)]
    samples = sanitize_sample_rows(rows, header_count=15)
    assert len(samples) == 3
    assert all(len(r) == 15 for r in samples)

    est = estimate_payload_tokens(clean, samples)
    assert est < 500


def test_preview_includes_mapping_source(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Consignee", "Grade", "Qty"])
    ws.append(["Z", "G1", 7])
    path = _save(wb, tmp_path / "preview_src.xlsx")

    with patch("app.erp_parser.parser_service.LLMHeaderResolver") as resolver_cls:
        preview = ERPParserService().preview(path)
        # High confidence known synonyms — LLM not required
        if preview["confidence"]["overall"] >= LLM_FALLBACK_THRESHOLD:
            resolver_cls.assert_not_called()

    assert "mapping_source" in preview
    assert preview["mapping_source"] in {"python", "llm"}


def test_scoring_preview_disables_llm(tmp_path: Path):
    """Background scoring path must not invoke LLM even on obscure headers."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Who Bought It", "Thing Sold SKU", "Bags Moved"])
    ws.append(["Omega Mills", "Latex-100", 55])
    path = _save(wb, tmp_path / "score_no_llm.xlsx")

    with patch("app.erp_parser.parser_service.LLMHeaderResolver") as resolver_cls:
        # May fail to map without LLM — that is OK for scoring; must not call OpenAI
        try:
            ERPParserService().preview(path, allow_llm_fallback=False)
        except Exception:
            pass
        resolver_cls.assert_not_called()


def test_explicit_preview_can_call_llm(tmp_path: Path):
    """Explicit Preview keeps LLM fallback enabled by default."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Who Bought It", "Thing Sold SKU", "Bags Moved"])
    ws.append(["Omega Mills Pvt", "Latex-100 Grade", 55])
    path = _save(wb, tmp_path / "preview_llm.xlsx")

    mock_resolver = MagicMock()
    mock_resolver.resolve_headers.return_value = {
        "customer_column": "Who Bought It",
        "product_column": "Thing Sold SKU",
        "quantity_column": "Bags Moved",
        "confidence": 96,
    }
    with patch(
        "app.erp_parser.parser_service.LLMHeaderResolver",
        return_value=mock_resolver,
    ):
        result = ERPParserService().preview(path, allow_llm_fallback=True)

    mock_resolver.resolve_headers.assert_called_once()
    assert result["mapping_source"] == "llm"
    assert abs(float(result["rows"][0]["sales_quantity"]) - 0.055) < 1e-9


def test_resolver_builds_openai_request_shape():
    """Unit: resolver posts temperature=0 JSON payload (mocked HTTP)."""
    captured = {}

    class FakeResponse:
        status_code = 200

        def raise_for_status(self):
            return None

        def json(self):
            return {
                "choices": [
                    {
                        "message": {
                            "content": (
                                '{"customer_column":"Sold To","product_column":"Commodity",'
                                '"quantity_column":"Net Movement","confidence":96}'
                            )
                        }
                    }
                ]
            }

    class FakeClient:
        def __init__(self, *a, **k):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def post(self, url, headers=None, json=None):
            captured["url"] = url
            captured["json"] = json
            captured["headers"] = headers
            return FakeResponse()

    with patch("app.erp_parser.llm_header_resolver.httpx.Client", FakeClient):
        out = LLMHeaderResolver(api_key="test-key", model="gpt-5.4-mini").resolve_headers(
            ["Sold To", "Commodity", "Net Movement", "Extra1"],
            [["A", "B", "1"], ["C", "D", "2"]],
        )

    assert out["customer_column"] == "Sold To"
    assert out["confidence"] == 96
    assert captured["json"]["temperature"] == 0
    assert captured["json"]["model"] == "gpt-5.4-mini"
    user = captured["json"]["messages"][1]["content"]
    assert "headers" in user
    assert "sample_rows" in user
    assert "workbook" not in user.lower()
