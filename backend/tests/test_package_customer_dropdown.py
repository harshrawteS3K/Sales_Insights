"""Package/draft Excel still includes Customer Name dropdown (Graph path mocked)."""

from pathlib import Path
from uuid import uuid4
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.config import settings
from app.database.session import SessionLocal
from app.integrations.graph.client import GraphClient
from app.repositories.distributor_repository import DistributorRepository
from app.services.product_master_service import ProductMasterService
from app.services.report_service import ReportService
from app.services.template_generation_service import TemplateGenerationService
from tests.workbook_helpers import build_official_workbook


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
        session.rollback()
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()


def _uid(prefix: str) -> str:
    return f"{prefix} {uuid4().hex[:8]}"


def _seed_products(db: Session) -> None:
    try:
        ProductMasterService(db).replace_from_rows(
            [("Paper", "P1"), ("Paper", "P2")], actor="test"
        )
    except Exception:
        from app.models.product_master import ProductMaster

        for industry, code in [("Paper", "P1"), ("Paper", "P2")]:
            db.add(
                ProductMaster(
                    industry_type=industry,
                    product_code=code,
                    is_active=True,
                )
            )
        db.flush()


def test_email_draft_excel_has_customer_and_product_dropdowns(db: Session, tmp_path: Path):
    _seed_products(db)
    company = _uid("Pkg Co")
    path = build_official_workbook(
        tmp_path / "in.xlsx",
        distributor="Harsh Rawte",
        company=company,
        reporting_month="Q2 2093",
        rows=[
            (1, "AKS RUGS", "Paper", "P1", 10),
            (2, "XYZ POLYMERS", "Paper", "P1", 5),
        ],
    )
    ReportService(db).ingest_excel(path, actor="tester")
    dist = DistributorRepository(db).get_by_company(company)
    assert dist is not None
    dist.email = f"{uuid4().hex[:8]}@example.com"
    db.flush()

    mock_graph = MagicMock(spec=GraphClient)
    mock_graph.resolve_mailbox.return_value = "salesinsights@apcotex.com"
    mock_graph.create_draft_message.return_value = {"id": "draft-xyz"}
    mock_graph.add_file_attachment.return_value = {"id": "att-1"}

    svc = TemplateGenerationService(db)
    svc.graph = mock_graph
    result = svc.create_email_draft(
        distributor_id=dist.id,
        reporting_quarter="Q3 2093",
        actor="tester",
    )

    excel_path = Path(settings.download_dir) / "packages" / result.attachment_name
    assert excel_path.is_file()
    wb = load_workbook(excel_path)
    assert "DistributorCustomers" in wb.defined_names
    attr = wb.defined_names["DistributorCustomers"].attr_text
    assert "_lists" in attr
    sheet = wb[wb.sheetnames[0]]
    formulas = [str(dv.formula1 or "") for dv in sheet.data_validations.dataValidation]
    assert any("DistributorCustomers" in f for f in formulas)
    assert any("SegmentList" in f for f in formulas)
    assert any("INDIRECT" in f for f in formulas)  # product dependent dropdown

    vals = [
        str(wb["CustomerList"].cell(row=r, column=1).value or "").strip().casefold()
        for r in range(2, wb["CustomerList"].max_row + 1)
    ]
    assert "aks rugs" in vals
    assert "xyz polymers" in vals
    mock_graph.create_draft_message.assert_called_once()
    assert mock_graph.sendMail.call_count == 0 if hasattr(mock_graph, "sendMail") else True
