"""Tests for Outlook link, validation summary, and viz consistency helpers."""

from pathlib import Path
from unittest.mock import MagicMock

import pytest
from sqlalchemy.orm import Session

from app.database.session import SessionLocal
from app.exceptions import GraphAPIError, NotFoundError
from app.models.email_message import EmailMessage
from app.repositories.email_repository import EmailMessageRepository
from app.services.dashboard_service import DashboardService
from app.services.outlook_sync_service import OutlookSyncService
from app.services.report_service import ReportService
from app.utils.datetime_utils import utc_now
from app.utils.outlook_links import resolve_outlook_open_url
from app.utils.validation_summary import build_validation_summary, validation_user_message
from tests.workbook_helpers import build_official_workbook


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
        session.rollback()
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()


def test_resolve_outlook_open_url_prefers_web_link():
    url = resolve_outlook_open_url(
        web_link="https://outlook.office.com/mail/deeplink/read/ABC",
        graph_message_id="gid",
    )
    assert url.startswith("https://outlook.office.com/")


def test_validation_summary_counts_reasons():
    summary = build_validation_summary(
        expected_rows=10,
        imported_rows=7,
        incomplete_rows=3,
        row_errors=[
            "Row 2 (excel row 11): Missing Customer",
            "Row 3 (excel row 12): Missing Product; Missing Quantity",
            "Row 4 (excel row 13): Missing Customer",
        ],
    )
    assert summary["skipped_rows"] == 3
    assert summary["imported_rows"] == 7
    reasons = {w["reason"]: w["count"] for w in summary["warnings"]}
    assert reasons["Missing Customer"] == 2
    assert reasons["Missing Product"] == 1
    msg = validation_user_message(summary)
    assert msg and "3 row" in msg


def test_ingest_persists_validation_summary(db: Session, tmp_path: Path):
    # Official workbook with one incomplete data row mixed with valid
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Name of Distributor"
    ws["B1"] = "Val Dist QA"
    ws["A2"] = "Company Name"
    ws["B2"] = "Co"
    ws["A3"] = "Address"
    ws["B3"] = "Addr"
    ws["A4"] = "Phone No"
    ws["B4"] = "1"
    ws["A5"] = "Reporting Month"
    ws["B5"] = "March 2099"
    headers = [
        "Sr. No.",
        "Name of Customer",
        "Segment",
        "Product",
        "Opening Stock",
        "Closing Stock",
        "Quantity",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(9, i, h)
    # valid
    for c, v in enumerate([1, "Cust A", "Paper", "P1", 1, 2, 10], 1):
        ws.cell(10, c, v)
    # missing customer + product
    for c, v in enumerate([2, "", "Paper", "", 1, 2, 5], 1):
        ws.cell(11, c, v)
    path = tmp_path / "val.xlsx"
    wb.save(path)

    report, inserted, _, _ = ReportService(db).ingest_excel(path, actor="test")
    assert inserted == 1
    extraction = (report.categories or {}).get("extraction") or {}
    summary = extraction.get("validation_summary") or {}
    assert summary.get("skipped_rows") == 1
    assert summary.get("imported_rows") == 1
    assert summary.get("warnings")


def test_outlook_link_404_raises_not_found(db: Session):
    email = EmailMessageRepository(db).create(
        EmailMessage(
            graph_message_id="gone-msg",
            subject="Gone",
            sender_email="a@b.com",
            received_at=utc_now(),
            mailbox="mb@test.com",
            outlook_web_link="https://outlook.office.com/mail/deeplink/read/gone",
            has_attachments=True,
        )
    )
    graph = MagicMock()
    graph.get_message = MagicMock(
        side_effect=GraphAPIError(
            "Microsoft Graph API returned HTTP 404",
            details='{"error":{"code":"ErrorItemNotFound"}}',
        )
    )
    svc = OutlookSyncService(db, graph_client=graph)
    with pytest.raises(NotFoundError) as exc:
        svc.resolve_outlook_open_link(email.id)
    assert "no longer available" in str(exc.value).lower()


def test_dashboard_distributor_count_matches_sales(db: Session, tmp_path: Path):
    path = build_official_workbook(
        tmp_path / "viz.xlsx",
        distributor="Viz Dist Only",
        reporting_month="April 2099",
        rows=[(1, "C1", "Paper", "P1", 10, 1, 2)],
    )
    ReportService(db).ingest_excel(path, actor="test")
    summary = DashboardService(db).summary()
    dist_rows = DashboardService(db).distributor_totals()
    # KPI distributor count must equal distributors that have active sales
    assert summary.total_distributors == len(dist_rows)
