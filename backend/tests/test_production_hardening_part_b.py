"""Part B production hardening regression tests (bugs 7–11)."""

from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.config import settings
from app.database.session import SessionLocal
from app.dependencies.rbac import get_current_user
from app.enums import UserRole
from app.exceptions import UnauthorizedError
from app.integrations.excel.template_generator import ExcelTemplateGenerator
from app.integrations.graph.client import GraphClient
from app.models.customer_master import CustomerMaster
from app.models.product_master import ProductMaster
from app.repositories.master_repository import CustomerMasterRepository, ProductMasterRepository
from app.services.report_service import ReportService
from app.enums import ReportSource
from app.utils.db_locks import report_replace_lock_key
from tests.workbook_helpers import build_official_workbook


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
        session.rollback()
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()


# ---------------------------------------------------------------------------
# BUG 7 — Outlook pagination
# ---------------------------------------------------------------------------


def test_bug7_graph_follows_odata_next_link():
    client = GraphClient(
        tenant_id="t",
        client_id="c",
        client_secret="s",
        mailbox="test@example.com",
        base_url="https://graph.example.com/v1.0",
    )
    page1 = {
        "value": [
            {"id": "m1", "subject": "One", "receivedDateTime": "2026-01-01T00:00:00Z"},
            {"id": "m2", "subject": "Two", "receivedDateTime": "2026-01-02T00:00:00Z"},
        ],
        "@odata.nextLink": "https://graph.example.com/v1.0/users/x/messages?$skiptoken=abc",
    }
    page2 = {
        "value": [
            {"id": "m3", "subject": "Three", "receivedDateTime": "2026-01-03T00:00:00Z"},
        ],
    }
    calls: list[tuple] = []

    def fake_request(method, path, params=None, **kwargs):
        calls.append((method, path, params))
        if "skiptoken" in path:
            return page2
        return page1

    with patch.object(client, "_request", side_effect=fake_request):
        with patch.object(client, "resolve_mailbox", return_value="test@example.com"):
            messages = client.list_unread_messages(top=50, page_size=2)

    assert [m["id"] for m in messages] == ["m1", "m2", "m3"]
    assert len(calls) == 2
    assert calls[0][2] is not None and "$top" in calls[0][2]
    assert calls[1][1].endswith("$skiptoken=abc")
    assert calls[1][2] is None  # nextLink must not re-apply params


def test_bug7_graph_respects_max_messages_cap():
    client = GraphClient(
        tenant_id="t",
        client_id="c",
        client_secret="s",
        mailbox="test@example.com",
        base_url="https://graph.example.com/v1.0",
    )
    page1 = {
        "value": [{"id": f"m{i}", "subject": str(i)} for i in range(3)],
        "@odata.nextLink": "https://graph.example.com/next",
    }

    with patch.object(client, "_request", return_value=page1):
        with patch.object(client, "resolve_mailbox", return_value="test@example.com"):
            messages = client.list_unread_messages(top=2, page_size=3)

    assert len(messages) == 2
    assert [m["id"] for m in messages] == ["m0", "m1"]


# ---------------------------------------------------------------------------
# BUG 8 — Report replace concurrency lock
# ---------------------------------------------------------------------------


def test_bug8_report_lock_keys_stable_for_business_identity():
    a = report_replace_lock_key(42, "July 2026")
    b = report_replace_lock_key(42, "  july   2026 ")
    c = report_replace_lock_key(43, "July 2026")
    assert a == b
    assert a != c


def test_bug8_sequential_replace_still_one_active(db: Session, tmp_path: Path):
    svc = ReportService(db)
    p1 = build_official_workbook(
        tmp_path / "lock1.xlsx",
        distributor="Lock Dist Co",
        reporting_month="August 2026",
        rows=[(1, "Cust A", "Paper", "P1", 10, 1, 2)],
    )
    p2 = build_official_workbook(
        tmp_path / "lock2.xlsx",
        distributor="Lock Dist Co",
        reporting_month="August 2026",
        rows=[(1, "Cust B", "Paper", "P2", 20, 3, 4)],
    )
    r1, _, _, _ = svc.ingest_excel(p1, source=ReportSource.UPLOAD, actor="admin-a")
    r2, _, _, _ = svc.ingest_excel(p2, source=ReportSource.UPLOAD, actor="admin-b")
    db.refresh(r1)
    db.refresh(r2)
    assert r1.is_deleted is True
    assert r2.is_deleted is False
    assert r1.distributor_id == r2.distributor_id


# ---------------------------------------------------------------------------
# BUG 9 — Master history retention purge
# ---------------------------------------------------------------------------


def test_bug9_customer_master_purges_old_soft_deleted(db: Session):
    repo = CustomerMasterRepository(db)
    old = CustomerMaster(
        customer_name="Ancient Customer",
        is_active=False,
        is_deleted=True,
        deleted_at=datetime.now(timezone.utc) - timedelta(days=200),
    )
    recent = CustomerMaster(
        customer_name="Recent Soft Deleted",
        is_active=False,
        is_deleted=True,
        deleted_at=datetime.now(timezone.utc) - timedelta(days=10),
    )
    db.add_all([old, recent])
    db.flush()
    old_id, recent_id = old.id, recent.id

    purged = repo.purge_soft_deleted_older_than(180)
    assert purged >= 1
    assert db.get(CustomerMaster, old_id) is None
    assert db.get(CustomerMaster, recent_id) is not None


def test_bug9_purge_disabled_when_days_zero(db: Session):
    repo = ProductMasterRepository(db)
    row = ProductMaster(
        industry_type="Paper",
        product_code="OLD-CODE",
        product_name="OLD-CODE",
        is_active=False,
        is_deleted=True,
        deleted_at=datetime.now(timezone.utc) - timedelta(days=400),
    )
    db.add(row)
    db.flush()
    rid = row.id
    assert repo.purge_soft_deleted_older_than(0) == 0
    assert db.get(ProductMaster, rid) is not None


# ---------------------------------------------------------------------------
# BUG 10 — Template scalability (identical layout)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("n", [500, 1000, 5000])
def test_bug10_template_scales(tmp_path: Path, n: int):
    products = [f"PROD-{i:05d}" for i in range(min(n, 5000))]
    out = tmp_path / f"template_{n}.xlsx"
    path = ExcelTemplateGenerator().generate(
        products=products,
        output_path=out,
    )
    assert path.exists()
    wb = load_workbook(path, read_only=True, data_only=False)
    assert "_lists" in wb.sheetnames
    lists = wb["_lists"]
    # Segments in column B; products under GENERAL for flat list
    assert lists["B1"].value == "Segments"
    assert lists["B2"].value == "GENERAL"
    assert "SegmentList" in wb.defined_names
    assert "CustomerList" not in wb.defined_names
    wb.close()


# ---------------------------------------------------------------------------
# BUG 11 — Auth modes
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_bug11_trusted_headers_rejects_missing_secret(monkeypatch):
    monkeypatch.setattr(settings, "auth_mode", "trusted_headers")
    monkeypatch.setattr(settings, "auth_trusted_secret", "corp-secret")
    monkeypatch.setattr(settings, "auth_trusted_header", "X-Internal-Auth")
    monkeypatch.setattr(settings, "app_env", "production")

    request = MagicMock()
    request.headers = {}

    with pytest.raises(UnauthorizedError):
        await get_current_user(request, x_user_role="admin", x_user_name="Admin", x_user_id=None)


@pytest.mark.asyncio
async def test_bug11_trusted_headers_accepts_valid_secret(monkeypatch):
    monkeypatch.setattr(settings, "auth_mode", "trusted_headers")
    monkeypatch.setattr(settings, "auth_trusted_secret", "corp-secret")
    monkeypatch.setattr(settings, "auth_trusted_header", "X-Internal-Auth")
    monkeypatch.setattr(settings, "app_env", "production")

    request = MagicMock()
    request.headers = {"X-Internal-Auth": "corp-secret"}
    request.state = MagicMock()

    user = await get_current_user(
        request, x_user_role="admin", x_user_name="Admin User", x_user_id="7"
    )
    assert user.role == UserRole.ADMIN
    assert user.name == "Admin User"
    assert user.user_id == 7


@pytest.mark.asyncio
async def test_bug11_headers_mode_preserves_dev_workflow(monkeypatch):
    monkeypatch.setattr(settings, "auth_mode", "headers")
    monkeypatch.setattr(settings, "app_env", "development")

    request = MagicMock()
    request.headers = {}
    request.state = MagicMock()

    user = await get_current_user(
        request, x_user_role="admin", x_user_name="Local Admin", x_user_id=None
    )
    assert user.role == UserRole.ADMIN
    assert user.name == "Local Admin"
