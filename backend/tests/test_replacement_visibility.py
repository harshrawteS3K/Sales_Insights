"""Tests for reporting-month normalization and replacement visibility."""

from datetime import datetime
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.database.session import SessionLocal
from app.models.report import Report
from app.models.sales_record import SalesRecord
from app.services.consolidated_data_service import ConsolidatedDataService
from app.services.report_service import ReportService
from app.utils.reporting_month import normalize_reporting_month
from tests.workbook_helpers import build_official_workbook


def test_normalize_reporting_month_from_datetime():
    assert normalize_reporting_month(datetime(2026, 7, 1)) == "July 2026"
    assert normalize_reporting_month("2026-07-01 00:00:00") == "July 2026"
    assert normalize_reporting_month("2026-07-01") == "July 2026"
    assert normalize_reporting_month("July 2026") == "July 2026"
    assert normalize_reporting_month("Q2 FY26") == "Q2 FY26"
    assert normalize_reporting_month(None) == ""


def test_excel_date_cell_becomes_july_2026(tmp_path: Path):
    from app.integrations.excel.parser import ExcelParserService

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Name of Distributor"
    ws["B1"] = "Date Dist"
    ws["A2"] = "Company Name"
    ws["B2"] = "Co"
    ws["A3"] = "Address"
    ws["B3"] = "Addr"
    ws["A4"] = "Phone No"
    ws["B4"] = "1"
    ws["A5"] = "Reporting Month"
    ws["B5"] = datetime(2026, 7, 1)  # Excel date cell
    for i, h in enumerate(
        ["Sr. No.", "Name of Customer", "Segment", "Product", "Quantity"], 1
    ):
        ws.cell(9, i, h)
    for c, v in enumerate([1, "Cust", "Carpet", "P1", 10], 1):
        ws.cell(10, c, v)
    path = tmp_path / "date_month.xlsx"
    wb.save(path)

    result = ExcelParserService().parse_sales_report(path)
    assert result.reporting_month == "July 2026"


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
        session.rollback()
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()


def test_replacement_hides_old_from_consolidated(db, tmp_path: Path):
    dist = f"Vis Dist {uuid4().hex[:8]}"
    month = "July 2098"
    v1 = build_official_workbook(
        tmp_path / "v1.xlsx",
        distributor=dist,
        reporting_month=month,
        include_stock=False,
        rows=[(1, "Old Cust", "Carpet", "P1", 50)],
    )
    v2 = build_official_workbook(
        tmp_path / "v2.xlsx",
        distributor=dist,
        reporting_month=month,
        include_stock=False,
        rows=[(1, "New Cust", "Carpet", "P2", 99)],
    )
    service = ReportService(db)
    old, _, _, _ = service.ingest_excel(v1, actor="test")
    new, _, _, _ = service.ingest_excel(v2, actor="test")
    db.flush()

    db.refresh(old)
    assert old.is_deleted is True
    assert new.is_deleted is False

    old_active = db.scalar(
        select(SalesRecord).where(
            SalesRecord.report_id == old.id,
            SalesRecord.is_deleted.is_(False),
        )
    )
    assert old_active is None

    page = ConsolidatedDataService(db).list_records(limit=5000, distributor=dist)
    assert len(page.data) == 1
    assert page.data[0].reportId == new.id
    assert page.data[0].sales[0].customerName == "New Cust"
    assert all(g.reportId != old.id for g in page.data)


def test_datetime_month_replaces_prior_text_month(db, tmp_path: Path):
    """Excel date cell for July must retire an existing 'July YYYY' active report."""
    dist = f"DateReplace {uuid4().hex[:8]}"
    text_path = build_official_workbook(
        tmp_path / "text.xlsx",
        distributor=dist,
        reporting_month="July 2097",
        include_stock=False,
        rows=[(1, "A", "Carpet", "P1", 10)],
    )
    service = ReportService(db)
    old, _, _, _ = service.ingest_excel(text_path, actor="test")

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Name of Distributor"
    ws["B1"] = dist
    ws["A2"] = "Company Name"
    ws["B2"] = "Co"
    ws["A3"] = "Address"
    ws["B3"] = "Addr"
    ws["A4"] = "Phone No"
    ws["B4"] = "1"
    ws["A5"] = "Reporting Month"
    ws["B5"] = datetime(2097, 7, 15)
    for i, h in enumerate(
        ["Sr. No.", "Name of Customer", "Segment", "Product", "Quantity"], 1
    ):
        ws.cell(9, i, h)
    for c, v in enumerate([1, "B", "Carpet", "P2", 20], 1):
        ws.cell(10, c, v)
    date_path = tmp_path / "date.xlsx"
    wb.save(date_path)

    new, inserted, _, _ = service.ingest_excel(date_path, actor="test")
    db.flush()
    assert inserted == 1
    assert new.reporting_month == "July 2097"
    db.refresh(old)
    assert old.is_deleted is True

    active = list(
        db.scalars(
            select(Report).where(
                Report.distributor_id == new.distributor_id,
                Report.is_deleted.is_(False),
            )
        ).all()
    )
    assert len(active) == 1
    assert active[0].id == new.id


def test_month_format_retires_legacy_quarter(db, tmp_path: Path):
    """Uploading July YYYY must archive an active Q2 FY## report for same distributor."""
    dist = f"LegacyQ {uuid4().hex[:8]}"
    q_path = build_official_workbook(
        tmp_path / "q.xlsx",
        distributor=dist,
        reporting_month="Q2 FY98",
        include_stock=False,
        rows=[(1, "Old", "Carpet", "P1", 10)],
    )
    m_path = build_official_workbook(
        tmp_path / "m.xlsx",
        distributor=dist,
        reporting_month="July 2098",
        include_stock=False,
        rows=[(1, "New", "Carpet", "P2", 20)],
    )
    service = ReportService(db)
    old, _, _, _ = service.ingest_excel(q_path, actor="test")
    new, _, _, _ = service.ingest_excel(m_path, actor="test")
    db.flush()
    db.refresh(old)
    assert old.is_deleted is True
    assert new.is_deleted is False
    assert new.reporting_month == "July 2098"

    page = ConsolidatedDataService(db).list_records(limit=5000, distributor=dist)
    assert len(page.data) == 1
    assert page.data[0].reportId == new.id
    assert page.data[0].reportingMonth == "July 2098"

