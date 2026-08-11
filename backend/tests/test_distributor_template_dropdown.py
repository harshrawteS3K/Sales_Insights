"""Ensure distributor-specific templates include Customer Name dropdown + correct download."""

from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.config import settings
from app.database.session import SessionLocal
from app.repositories.distributor_repository import DistributorRepository
from app.services.product_master_service import ProductMasterService
from app.services.report_service import ReportService
from app.services.template_generation_service import TemplateGenerationService
from tests.workbook_helpers import build_official_workbook


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
        session.rollback()
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()


def _uid(prefix: str) -> str:
    return f"{prefix} {uuid4().hex[:8]}"


def _seed_products(db: Session) -> None:
    try:
        ProductMasterService(db).replace_from_rows(
            [("Paper", "P1"), ("Paper", "P2")], actor="test"
        )
    except Exception:
        from app.models.product_master import ProductMaster

        for industry, code in [("Paper", "P1"), ("Paper", "P2")]:
            db.add(
                ProductMaster(
                    industry_type=industry,
                    product_code=code,
                    is_active=True,
                )
            )
        db.flush()


def test_distributor_template_has_customer_dropdown(db: Session, tmp_path: Path):
    _seed_products(db)
    company = _uid("Demo Drop")
    path = build_official_workbook(
        tmp_path / "in.xlsx",
        distributor="Harsh Rawte",
        company=company,
        reporting_month="Q2 2093",
        rows=[(1, "AKS RUGS", "Paper", "P1", 10)],
    )
    ReportService(db).ingest_excel(path, actor="tester")
    dist = DistributorRepository(db).get_by_company(company)
    assert dist is not None

    out = Path(settings.download_dir) / f"Apcotex_Test_Dropdown_{uuid4().hex[:6]}.xlsx"
    result = TemplateGenerationService(db).generate(
        actor="tester",
        mode="distributor",
        distributor_id=dist.id,
        reporting_quarter=None,
        output_path=out,
        update_canonical=False,
    )
    assert result.fallback_generic is False
    assert result.customers_count >= 1
    assert out.is_file()

    wb = load_workbook(out)
    assert "CustomerList" in wb.sheetnames
    assert "DistributorCustomers" in wb.defined_names

    sheet = wb[wb.sheetnames[0]]
    formulas = [str(dv.formula1 or "") for dv in sheet.data_validations.dataValidation]
    assert any("DistributorCustomers" in f for f in formulas)

    # Customer Name column B data rows must be covered by a list validation
    ranges = []
    for dv in sheet.data_validations.dataValidation:
        if "DistributorCustomers" in str(dv.formula1 or ""):
            ranges.extend(list(dv.sqref))
    joined = " ".join(str(r) for r in ranges)
    assert "B7" in joined or "B7:B1000" in joined or any(
        "B7" in str(r) for r in ranges
    )

    cust_vals = [
        str(wb["CustomerList"].cell(row=r, column=1).value or "").strip()
        for r in range(2, wb["CustomerList"].max_row + 1)
    ]
    assert any(v.casefold() == "aks rugs" for v in cust_vals)

    # Download resolver must return THIS file when preferred_name is set
    resolved, name = TemplateGenerationService(db).resolve_download_path(
        preferred_name=out.name
    )
    assert resolved.resolve() == out.resolve()
    assert name == out.name

    # Without preferred_name, newest Apcotex_*.xlsx should still be findable
    resolved2, _ = TemplateGenerationService(db).resolve_download_path()
    assert resolved2.is_file()
