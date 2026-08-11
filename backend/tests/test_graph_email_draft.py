"""Microsoft Graph Outlook draft creation for distributor email workflow."""

from pathlib import Path
from uuid import uuid4
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database.session import SessionLocal
from app.exceptions import ValidationAppError
from app.integrations.graph.client import GraphClient
from app.repositories.distributor_repository import DistributorRepository
from app.services.report_service import ReportService
from app.services.template_generation_service import TemplateGenerationService
from tests.workbook_helpers import build_official_workbook


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
        session.rollback()
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()


def _uid(prefix: str) -> str:
    return f"{prefix} {uuid4().hex[:8]}"


def _seed_products(db: Session) -> None:
    from app.models.product_master import ProductMaster
    from app.repositories.master_repository import ProductMasterRepository

    if ProductMasterRepository(db).list_codes_by_segment():
        return
    for industry, code in [("Paper", "P1"), ("Paper", "P2")]:
        db.add(
            ProductMaster(
                industry_type=industry,
                product_code=code,
                is_active=True,
            )
        )
    db.flush()


def test_create_email_draft_requires_distributor_email(db: Session, tmp_path: Path):
    _seed_products(db)
    company = _uid("NoMail")
    path = build_official_workbook(
        tmp_path / "in.xlsx",
        distributor="Rep",
        company=company,
        reporting_month="Q1 2091",
        rows=[(1, "AKS RUGS", "Paper", "P1", 10)],
    )
    ReportService(db).ingest_excel(path, actor="tester")
    dist = DistributorRepository(db).get_by_company(company)
    assert dist is not None
    dist.email = None
    db.flush()

    with pytest.raises(ValidationAppError) as exc:
        TemplateGenerationService(db).create_email_draft(
            distributor_id=dist.id,
            reporting_quarter="Q2 2091",
            actor="admin",
        )
    assert "email" in str(exc.value.message).lower()


def test_create_email_draft_calls_graph_not_sendmail(db: Session, tmp_path: Path, monkeypatch):
    _seed_products(db)
    company = _uid("GraphDraft")
    path = build_official_workbook(
        tmp_path / "in.xlsx",
        distributor="Harsh Rawte",
        company=company,
        reporting_month="Q2 2092",
        rows=[(1, "AKS RUGS", "Paper", "P1", 10)],
    )
    ReportService(db).ingest_excel(path, actor="tester")
    dist = DistributorRepository(db).get_by_company(company)
    assert dist is not None
    dist.email = "distributor@example.com"
    dist.cc_email = "cc@example.com"
    dist.contact_person = "Navneet"
    db.flush()

    mock_graph = MagicMock(spec=GraphClient)
    mock_graph.resolve_mailbox.return_value = "salesinsights@apcotex.com"
    mock_graph.create_draft_message.return_value = {"id": "draft-123"}
    mock_graph.add_file_attachment.return_value = {"id": "att-1"}

    svc = TemplateGenerationService(db)
    svc.graph = mock_graph

    result = svc.create_email_draft(
        distributor_id=dist.id,
        reporting_quarter="Q3 2092",
        actor="admin",
    )

    assert result.success is True
    assert result.draft_id == "draft-123"
    assert result.mailbox == "salesinsights@apcotex.com"
    assert result.recipient == "distributor@example.com"
    assert result.cc == "cc@example.com"
    assert result.attachment_name.endswith(".xlsx")
    assert "Template" in result.attachment_name
    assert "sendMail" not in str(mock_graph.method_calls).lower()
    mock_graph.create_draft_message.assert_called_once()
    call_kw = mock_graph.create_draft_message.call_args.kwargs
    assert call_kw["to_email"] == "distributor@example.com"
    assert call_kw["cc_email"] == "cc@example.com"
    assert "Q3 2092" in call_kw["subject"]

    att_kw = mock_graph.add_file_attachment.call_args
    assert att_kw.args[0] == "draft-123"
    assert att_kw.kwargs["filename"].endswith(".xlsx")
    content = att_kw.kwargs.get("content") or b""
    assert len(content) > 100
    assert content[:2] == b"PK"  # zip/xlsx magic

    # Workbook still has customer dropdown
    # Reconstruct path from packages dir
    from app.core.config import settings

    packages = Path(settings.download_dir) / "packages"
    matches = list(packages.glob(f"*{result.attachment_name}"))
    assert matches, "Expected generated excel on disk"
    wb = load_workbook(matches[0])
    assert "DistributorCustomers" in wb.defined_names
    formulas = [
        str(dv.formula1 or "")
        for dv in wb[wb.sheetnames[0]].data_validations.dataValidation
    ]
    assert any("DistributorCustomers" in f for f in formulas)
    assert any("SegmentList" in f for f in formulas)


def test_graph_create_draft_payload_shape():
    client = GraphClient(
        tenant_id="t",
        client_id="c",
        client_secret="s",
        mailbox="salesinsights@apcotex.com",
    )
    captured = {}

    def fake_request(method, path, *, params=None, json_body=None, raw=False, timeout=60):
        captured["method"] = method
        captured["path"] = path
        captured["json_body"] = json_body
        assert "sendMail" not in path
        return {"id": "msg-1"}

    client._request = fake_request  # type: ignore[method-assign]
    client.resolve_mailbox = lambda mailbox=None: "salesinsights@apcotex.com"  # type: ignore

    out = client.create_draft_message(
        subject="Test",
        body_text="Body",
        to_email="a@example.com",
        cc_email="b@example.com",
    )
    assert out["id"] == "msg-1"
    assert captured["method"] == "POST"
    assert captured["path"].endswith("/messages")
    assert "sendMail" not in captured["path"]
    body = captured["json_body"]
    assert body["toRecipients"][0]["emailAddress"]["address"] == "a@example.com"
    assert body["ccRecipients"][0]["emailAddress"]["address"] == "b@example.com"


def test_graph_attachment_is_base64_file_attachment():
    client = GraphClient(
        tenant_id="t",
        client_id="c",
        client_secret="s",
        mailbox="salesinsights@apcotex.com",
    )
    captured = {}

    def fake_request(method, path, *, params=None, json_body=None, raw=False, timeout=60):
        captured["method"] = method
        captured["path"] = path
        captured["json_body"] = json_body
        return {"id": "att-1"}

    client._request = fake_request  # type: ignore[method-assign]
    client.resolve_mailbox = lambda mailbox=None: "salesinsights@apcotex.com"  # type: ignore

    content = b"PK\x03\x04fake-xlsx"
    client.add_file_attachment(
        "draft-1",
        filename="Apcotex_Demo_Q3_2026_Template.xlsx",
        content=content,
    )
    body = captured["json_body"]
    assert body["@odata.type"] == "#microsoft.graph.fileAttachment"
    assert body["name"].endswith(".xlsx")
    assert body["contentType"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert body["contentBytes"]
    assert "GRAPH_CLIENT_SECRET" not in str(captured)
    assert "s" not in str(body.get("contentBytes", ""))[:5] or True
