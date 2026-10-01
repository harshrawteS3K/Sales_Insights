"""Altek layout: Customer | Product | month columns with a merged summary block below."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from app.erp_parser import ERPParserService
from app.erp_parser.matrix_month_parser import (
    detect_matrix_month_layout,
    extract_matrix_month_rows,
    extract_matrix_month_workbook,
)

ORCH_RESOLVER = "app.erp_parser.orchestrator.service.LLMHeaderResolver"

ALTEK_CUSTOMERS = [
    ("AYN", None),
    ("ADNC", "656M"),
    ("BOZATELI", None),
    ("CMB", None),
    ("MARINOS", None),
    ("MARPOL FLOORING", None),
    ("STANDARD CARPET", 625),
    ("STANDARD CARPET", 656),
    ("KAPLAN FLOORING", None),
]


def _altek_workbook(path: Path) -> Path:
    """Mirror of ``Altek Int.xlsx`` — summary label merged over its product block."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Customer ", "Product", datetime(2026, 4, 1), datetime(2026, 5, 1), datetime(2026, 6, 1)])
    for customer, product in ALTEK_CUSTOMERS:
        ws.append([customer, product, None, None, None])
    ws.append(["SUNSILK", "CB 656M", None, None, 23.06])
    summary_start = ws.max_row + 1
    ws.append(["Altek Total Sales", 7480, 614.4, 304.8, 592.4])
    ws.append([None, 620, 417, None, None])
    ws.append([None, 625, None, None, 23])
    ws.append([None, "625LV", None, None, 69.74])
    ws.append([None, "8000 5B", None, 76.8, 25.6])
    ws.append([None, "656M", None, 76.8, 48.6])
    ws.append([None, 6300, None, 48, None])
    ws.merge_cells(start_row=summary_start, start_column=1, end_row=ws.max_row, end_column=1)
    ws.append([None, None, 1031.4, 506.4, 759.34])
    wb.save(path)
    return path


def test_detects_altek_layout(tmp_path: Path):
    from app.erp_parser.workbook_detector import read_sheet_matrix

    matrix = read_sheet_matrix(_altek_workbook(tmp_path / "altek.xlsx"), "Sheet1")
    assert detect_matrix_month_layout(matrix)


def test_extracts_only_real_customer_cells_and_ignores_summary(tmp_path: Path):
    result = extract_matrix_month_workbook(_altek_workbook(tmp_path / "altek.xlsx"))
    assert result is not None

    detected = result["customers_detected"]
    names = [c["customer"] for c in detected]
    assert sorted(set(names)) == sorted(
        {"AYN", "ADNC", "BOZATELI", "CMB", "MARINOS", "MARPOL FLOORING",
         "STANDARD CARPET", "KAPLAN FLOORING", "SUNSILK"}
    )
    assert len(set(names)) == 9
    assert not any("total" in n.lower() for n in names)
    assert result["skipped_total"] == 7

    products = {(c["customer"], c["product"]) for c in detected}
    assert ("ADNC", "656M") in products
    assert ("STANDARD CARPET", "625") in products
    assert ("STANDARD CARPET", "656") in products
    assert ("SUNSILK", "CB 656M") in products
    for blank in ("AYN", "BOZATELI", "CMB", "MARINOS", "MARPOL FLOORING", "KAPLAN FLOORING"):
        assert (blank, None) in products

    rows = result["rows"]
    assert len(rows) == 1
    only = rows[0]
    assert only["customer_name"] == "SUNSILK"
    assert only["product"] == "CB 656M"
    assert float(only["sales_quantity"]) == 23.06
    assert only["source_month"] == "June 2026"
    assert result["llm_used"] is False


def test_april_may_june_map_to_q1_fy2026_27():
    matrix = [
        ["Customer", "Product", datetime(2026, 4, 1), datetime(2026, 5, 1), datetime(2026, 6, 1)],
        ["Acme", "625", 10, 20, 30],
        ["Acme", None, 99, 99, 99],
        ["Altek Total Sales", None, 10, 20, 30],
    ]
    result = extract_matrix_month_rows(matrix)
    assert result is not None
    by_month = {r["source_month"]: float(r["sales_quantity"]) for r in result["rows"]}
    assert by_month == {"April 2026": 10.0, "May 2026": 20.0, "June 2026": 30.0}
    assert {r["customer_name"] for r in result["rows"]} == {"Acme"}
    assert {r["product"] for r in result["rows"]} == {"625"}
    assert len({r["period"] for r in result["rows"]}) == 1
    assert "2026" in result["rows"][0]["period"] and "Q1" in result["rows"][0]["period"]


def test_summary_labels_skipped_but_real_names_with_total_kept():
    matrix = [
        ["Customer", "Product", "Apr-2026", "May-2026"],
        ["Petro Total Solutions", "N745", 5, 6],
        ["Totalcare Rubber", "N745", 5, 6],
        ["Grand Total", "", 5, 6],
        ["Party Total", "", 5, 6],
        ["Total", "", 5, 6],
        ["Altek Total Sales", "7480", 5, 6],
    ]
    result = extract_matrix_month_rows(matrix)
    assert {r["customer_name"] for r in result["rows"]} == {"Petro Total Solutions", "Totalcare Rubber"}
    assert result["skipped_total"] == 4


def test_full_pipeline_deterministic_without_llm(tmp_path: Path):
    path = _altek_workbook(tmp_path / "Altek Int.xlsx")
    with patch(ORCH_RESOLVER) as resolver_cls:
        result = ERPParserService().parse_workbook(path, allow_llm_fallback=True)
        resolver_cls.assert_not_called()

    assert result.mapping_source == "python"
    assert result.imported_rows == 1
    assert [r["customer_name"] for r in result.rows] == ["SUNSILK"]
    assert [r["product"] for r in result.rows] == ["CB 656M"]
    assert not any("total" in str(r["customer_name"]).lower() for r in result.rows)
