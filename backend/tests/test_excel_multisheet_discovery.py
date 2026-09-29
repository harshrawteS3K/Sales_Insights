"""Multi-sheet Excel discovery: scan all sheets, pick the best sales table."""

import time
from pathlib import Path

from openpyxl import Workbook

from app.erp_parser import ERPParserService
from app.erp_parser.header_dictionary import get_header_dictionary
from app.erp_parser.header_mapper import map_headers
from app.erp_parser.workbook_detector import list_candidate_sheets


def _save(wb: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _irrelevant(ws) -> None:
    ws.append(["Cover page"])
    ws.append(["See other tabs for details"])


def _sales_table(ws, *, title: str | None = None) -> None:
    if title:
        ws.append([title])
        ws.append([])
    ws.append(["Party Name", "Item Name", "Actual Quantity"])
    ws.append(["PIONEER RUBBER INDUSTRIES PVT.LTD.", "N 745", 175.0])
    ws.append(["PUJA FLUID SEALS PVT.LTD.", "N 745", 315.0])
    ws.append(["SAMPLE SEALS CO.", "N 710", 50.5])


def test_a_three_sheet_discovers_sales_on_any_sheet(tmp_path: Path):
    """TEST A: 3-sheet workbook — sales data on any sheet must be found."""
    wb = Workbook()
    cover = wb.active
    cover.title = "Sheet1"
    _irrelevant(cover)
    mid = wb.create_sheet("Notes")
    _irrelevant(mid)
    sales = wb.create_sheet("Report")
    _sales_table(sales, title="Nitrile Report - June-2026")
    path = _save(wb, tmp_path / "three_sheet.xlsx")

    started = time.perf_counter()
    result = ERPParserService().parse_workbook(path, allow_llm_fallback=False)
    elapsed = time.perf_counter() - started

    assert set(result.candidate_sheets) >= {"Sheet1", "Notes", "Report"}
    assert result.sheet_name == "Report"
    assert result.imported_rows == 3
    assert result.rows[0]["customer_name"] == "PIONEER RUBBER INDUSTRIES PVT.LTD."
    assert result.rows[0]["product"] == "N 745"
    assert float(result.rows[0]["sales_quantity"]) == 175.0
    assert result.overall_confidence >= 75
    assert elapsed < 90.0
    print(
        f"A sheets={result.candidate_sheets} selected={result.sheet_name} "
        f"layout={(result.confidence_breakdown or {}).get('layout')} "
        f"rows={result.imported_rows} confidence={result.overall_confidence} time={elapsed:.2f}s"
    )


def test_b_sheet2_named_data_selected(tmp_path: Path):
    """TEST B: Sheet1 irrelevant, Sheet2 'Data' has the table → select Data."""
    wb = Workbook()
    first = wb.active
    first.title = "Sheet1"
    _irrelevant(first)
    data = wb.create_sheet("Data")
    data.append(["Customer", "Product", "Qty"])
    data.append(["Alpha Traders", "TX 400", 100])
    data.append(["Beta Corp", "SR 568", 200])
    path = _save(wb, tmp_path / "ab_brothers_style.xlsx")

    started = time.perf_counter()
    result = ERPParserService().parse_workbook(path, allow_llm_fallback=False)
    elapsed = time.perf_counter() - started

    assert "Data" in list_candidate_sheets(path)
    assert result.sheet_name == "Data"
    assert result.imported_rows == 2
    assert result.overall_confidence >= 75
    print(
        f"B sheets={result.candidate_sheets} selected={result.sheet_name} "
        f"rows={result.imported_rows} confidence={result.overall_confidence} time={elapsed:.2f}s"
    )


def test_c_avik_style_party_item_actual_quantity_mapping(tmp_path: Path):
    """TEST C: Party Name | Item Name | Actual Quantity maps Customer/Product/Qty."""
    get_header_dictionary().reload()
    mapped = map_headers(["Party Name", "Item Name", "Actual Quantity"])
    assert mapped["positions"]["customer"] == 0
    assert mapped["positions"]["product"] == 1
    assert mapped["positions"]["quantity"] == 2

    wb = Workbook()
    ws = wb.active
    ws.title = "June"
    _sales_table(ws, title="Nitrile Report - June-2026")
    path = _save(wb, tmp_path / "avik_style.xlsx")

    started = time.perf_counter()
    result = ERPParserService().parse_workbook(path, allow_llm_fallback=False)
    elapsed = time.perf_counter() - started

    assert result.imported_rows == 3
    assert result.rows[1]["customer_name"] == "PUJA FLUID SEALS PVT.LTD."
    assert result.rows[1]["product"] == "N 745"
    assert float(result.rows[1]["sales_quantity"]) == 315.0
    assert result.overall_confidence >= 75
    print(
        f"C selected={result.sheet_name} layout={(result.confidence_breakdown or {}).get('layout')} "
        f"rows={result.imported_rows} confidence={result.overall_confidence} time={elapsed:.2f}s"
    )


def test_d_irrelevant_sheets_do_not_fail_workbook(tmp_path: Path):
    """TEST D: Multiple irrelevant sheets + one valid → must not fail early."""
    wb = Workbook()
    a = wb.active
    a.title = "Intro"
    _irrelevant(a)
    b = wb.create_sheet("Summary")
    b.append(["Totals only"])
    b.append(["Grand Total", 9999])
    c = wb.create_sheet("SalesData")
    _sales_table(c)
    path = _save(wb, tmp_path / "noise_then_valid.xlsx")

    result = ERPParserService().parse_workbook(path, allow_llm_fallback=False)
    assert result.sheet_name == "SalesData"
    assert result.imported_rows == 3


def test_e_valid_data_on_sheet3_discovered(tmp_path: Path):
    """TEST E: Valid data only on Sheet3 must still be discovered."""
    wb = Workbook()
    s1 = wb.active
    s1.title = "Sheet1"
    _irrelevant(s1)
    s2 = wb.create_sheet("Sheet2")
    _irrelevant(s2)
    s3 = wb.create_sheet("Sheet3")
    s3.append(["Buyer", "Material", "Dispatch Qty"])
    s3.append(["Cust One", "FG-1", 10])
    s3.append(["Cust Two", "FG-2", 20])
    path = _save(wb, tmp_path / "sheet3_wins.xlsx")

    started = time.perf_counter()
    result = ERPParserService().parse_workbook(path, allow_llm_fallback=False)
    elapsed = time.perf_counter() - started

    assert result.sheet_name == "Sheet3"
    assert result.imported_rows == 2
    assert float(result.rows[0]["sales_quantity"]) == 10
    assert result.overall_confidence >= 70
    print(
        f"E sheets={result.candidate_sheets} selected={result.sheet_name} "
        f"rows={result.imported_rows} confidence={result.overall_confidence} time={elapsed:.2f}s"
    )
