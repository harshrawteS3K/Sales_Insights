"""Phase 2 Master Data Engine regression tests."""

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.database.session import SessionLocal
from app.exceptions import ExcelProcessingError, ValidationAppError
from app.integrations.excel.parser import ExcelParserService
from app.integrations.excel.template_generator import ExcelTemplateGenerator
from app.repositories.master_repository import CustomerMasterRepository, ProductMasterRepository
from app.services.customer_master_service import CustomerMasterService
from app.services.product_master_service import ProductMasterService
from app.services.template_generation_service import TemplateGenerationService
from app.validators.excel_validators import (
    validate_customer_master_dataframe,
    validate_product_master_dataframe,
)
import pandas as pd


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
        session.rollback()
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()


def _write_customer_master(path: Path, names: list[str], *, header: str = "CUSTOMER NAME") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append([header, "REGION"])
    for name in names:
        ws.append([name, "West"])
    wb.save(path)
    return path


def _write_product_master(
    path: Path,
    rows: list[tuple[str, str]],
    *,
    industry_header: str = "Industry Type Description",
    code_header: str = "Product Code",
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append([industry_header, code_header])
    for industry, code in rows:
        ws.append([industry, code])
    wb.save(path)
    return path


class _Upload:
    """Minimal UploadFile-like object for service tests."""

    def __init__(self, path: Path) -> None:
        self.filename = path.name
        self._path = path
        self._fh = None

    async def read(self, size: int = -1) -> bytes:
        if self._fh is None:
            self._fh = self._path.open("rb")
        if size < 0:
            data = self._fh.read()
        else:
            data = self._fh.read(size)
        return data

    async def seek(self, offset: int) -> None:
        if self._fh is None:
            self._fh = self._path.open("rb")
        self._fh.seek(offset)


@pytest.mark.asyncio
async def test_scenario1_upload_customer_master_1000(db: Session, tmp_path: Path):
    names = [f"Customer {i:04d}" for i in range(1000)]
    path = _write_customer_master(tmp_path / "customers.xlsx", names)
    svc = CustomerMasterService(db)
    result = await svc.upload_and_replace(_Upload(path), actor="test")
    assert result.success
    assert result.records_imported == 1000
    assert CustomerMasterRepository(db).count() == 1000


@pytest.mark.asyncio
async def test_scenario2_upload_product_master_200(db: Session, tmp_path: Path):
    rows = [(f"INDUSTRY {i % 10}", f"PROD-{i:04d}") for i in range(200)]
    path = _write_product_master(tmp_path / "products.xlsx", rows)
    svc = ProductMasterService(db)
    result = await svc.upload_and_replace(_Upload(path), actor="test")
    assert result.success
    assert result.records_imported == 200
    assert ProductMasterRepository(db).count() == 200


@pytest.mark.asyncio
async def test_scenario3_generate_template_dropdowns(db: Session, tmp_path: Path):
    cust_path = _write_customer_master(tmp_path / "c.xlsx", ["Alpha Co", "Beta Co", "Zeta Co"])
    prod_path = _write_product_master(
        tmp_path / "p.xlsx",
        [("PAPER", "FGC8200-200"), ("CARPET", "FGCB200-200")],
    )
    await CustomerMasterService(db).upload_and_replace(_Upload(cust_path), actor="test")
    await ProductMasterService(db).upload_and_replace(_Upload(prod_path), actor="test")

    out = tmp_path / "template.xlsx"
    ExcelTemplateGenerator().generate(
        customers=CustomerMasterRepository(db).list_names(),
        products_by_segment=ProductMasterRepository(db).list_codes_by_segment(),
        output_path=out,
    )
    wb = load_workbook(out)
    assert "Sales Report" in wb.sheetnames
    assert "_lists" in wb.sheetnames
    lists = wb["_lists"]
    assert lists.sheet_state == "hidden"
    customers = [lists["A2"].value, lists["A3"].value, lists["A4"].value]
    assert customers == ["Alpha Co", "Beta Co", "Zeta Co"]
    assert set(filter(None, [lists["B2"].value, lists["B3"].value])) == {"CARPET", "PAPER"}
    assert "SEG_CARPET" in wb.defined_names
    assert "SEG_PAPER" in wb.defined_names
    assert "SegmentMap" in wb.defined_names
    # Headers include Customer + Product for dropdowns
    sheet = wb["Sales Report"]
    headers = [sheet.cell(7, c).value for c in range(1, 8)]
    assert headers == [
        "Sr. No.",
        "Name of Customer",
        "Segment",
        "Product",
        "Opening Stock",
        "Closing Stock",
        "Quantity",
    ]
    formulas = [str(dv.formula1) for dv in sheet.data_validations.dataValidation]
    assert any("INDIRECT" in f and "SegmentMap" in f for f in formulas)
    assert len(sheet.data_validations.dataValidation) >= 2
    # Distributor header values must always be blank
    for row in range(1, 6):
        assert sheet.cell(row, 2).value in (None, "")


@pytest.mark.asyncio
async def test_generated_template_never_prefills_distributor_header(
    db: Session, tmp_path: Path
):
    """
    Regression: template must stay blank even when distributors / periods exist in DB
    and even if a distributors list is passed into the generator.
    """
    from app.models.distributor import Distributor
    from app.repositories.distributor_repository import DistributorRepository

    cust_path = _write_customer_master(tmp_path / "c.xlsx", ["Cust A"])
    prod_path = _write_product_master(tmp_path / "p.xlsx", [("PAPER", "PROD-1")])
    await CustomerMasterService(db).upload_and_replace(_Upload(cust_path), actor="test")
    await ProductMasterService(db).upload_and_replace(_Upload(prod_path), actor="test")

    # Seed a distributor that previously would have been written into B1
    DistributorRepository(db).create(
        Distributor(name="Harsh", company="Harsh Co", address="Pune", phone="999")
    )
    db.flush()

    out = tmp_path / "blank_template.xlsx"
    ExcelTemplateGenerator().generate(
        customers=["Cust A"],
        products=["PROD-1"],
        distributors=["Harsh", "Someone Else"],
        periods=["August 2026", "July 2026"],
        output_path=out,
    )
    sheet = load_workbook(out)["Sales Report"]
    assert sheet.cell(1, 1).value == "Name of Distributor"
    assert sheet.cell(2, 1).value == "Company Name"
    assert sheet.cell(3, 1).value == "Address"
    assert sheet.cell(4, 1).value == "Phone No"
    assert sheet.cell(5, 1).value == "Reporting Month"
    for row in range(1, 6):
        assert sheet.cell(row, 2).value in (None, ""), (
            f"Row {row} col B must be blank, got {sheet.cell(row, 2).value!r}"
        )

    # Service path must also produce a blank header
    result = TemplateGenerationService(db).generate(actor="test")
    assert result.success
    path, _ = TemplateGenerationService(db).resolve_download_path(
        preferred_name=result.file_name
    )
    svc_sheet = load_workbook(path)["Sales Report"]
    for row in range(1, 6):
        assert svc_sheet.cell(row, 2).value in (None, "")
    # Sales entry cells (except Sr. No.) stay blank
    for col in range(2, 8):
        assert svc_sheet.cell(8, col).value in (None, "")

    svc = CustomerMasterService(db)
    await svc.upload_and_replace(
        _Upload(_write_customer_master(tmp_path / "c1.xlsx", ["Old A", "Old B"])),
        actor="test",
    )
    assert set(CustomerMasterRepository(db).list_names()) == {"Old A", "Old B"}

    await svc.upload_and_replace(
        _Upload(_write_customer_master(tmp_path / "c2.xlsx", ["New X", "New Y", "New Z"])),
        actor="test",
    )
    names = CustomerMasterRepository(db).list_names()
    assert set(names) == {"New X", "New Y", "New Z"}
    assert CustomerMasterRepository(db).count() == 3


@pytest.mark.asyncio
async def test_scenario5_product_master_replace(db: Session, tmp_path: Path):
    svc = ProductMasterService(db)
    await svc.upload_and_replace(
        _Upload(_write_product_master(tmp_path / "p1.xlsx", [("PAPER", "OLD-1"), ("PAPER", "OLD-2")])),
        actor="test",
    )
    await svc.upload_and_replace(
        _Upload(
            _write_product_master(
                tmp_path / "p2.xlsx",
                [("CARPET", "NEW-A"), ("CONSTRUCTION B2B IH", "NEW-B")],
            )
        ),
        actor="test",
    )
    codes = ProductMasterRepository(db).list_product_codes()
    assert set(codes) == {"NEW-A", "NEW-B"}
    assert ProductMasterRepository(db).count() == 2


def test_scenario6_invalid_excel_extension(tmp_path: Path):
    from app.services.excel_validation_service import ExcelValidationService

    with pytest.raises(ExcelProcessingError):
        ExcelValidationService().assert_excel_file("master.csv")


def test_scenario7_missing_required_header(tmp_path: Path):
    # Customer master without CUSTOMER NAME
    bad_c = tmp_path / "bad_c.xlsx"
    wb = Workbook()
    wb.active.append(["Region", "City"])
    wb.active.append(["West", "Pune"])
    wb.save(bad_c)
    df = pd.read_excel(bad_c)
    with pytest.raises(ExcelProcessingError) as exc:
        validate_customer_master_dataframe(df)
    assert "missing required columns" in str(exc.value).lower() or "customer_name" in str(exc.value)

    # Product master missing Product Code
    bad_p = tmp_path / "bad_p.xlsx"
    wb = Workbook()
    wb.active.append(["Industry Type Description"])
    wb.active.append(["PAPER"])
    wb.save(bad_p)
    dfp = pd.read_excel(bad_p)
    with pytest.raises(ExcelProcessingError):
        validate_product_master_dataframe(dfp)


def test_parser_dedupes_and_trims(tmp_path: Path):
    path = _write_customer_master(
        tmp_path / "dup.xlsx",
        ["  Acme  ", "Acme", "Beta", "", "beta"],
    )
    parsed = ExcelParserService().parse_customer_master(path)
    names = [r["customer_name"] for r in parsed.records]
    assert names == ["Acme", "Beta"]
    assert parsed.duplicate_names == 2
    assert parsed.blank_customer_name == 1
    assert parsed.imported == 2
    assert parsed.excel_rows == 5
    assert parsed.skipped == 3
    assert "Blank Customer Name" in parsed.summary_message()


@pytest.mark.asyncio
async def test_customer_replace_returns_full_summary(db: Session, tmp_path: Path):
    path = _write_customer_master(
        tmp_path / "sum.xlsx",
        ["Alpha", "alpha", "Beta", "", "Gamma"],
    )
    result = await CustomerMasterService(db).upload_and_replace(_Upload(path), actor="test")
    assert result.records_imported == 3
    assert result.duplicate_names == 1
    assert result.blank_customer_name == 1
    assert result.excel_rows == 5
    assert result.records_skipped == 2
    assert "Customer Master Replace Complete" in result.message
    assert "Imported            : 3" in result.message


def test_segment_product_filter_excludes_other_segments(tmp_path: Path):
    """CARPET named range must not contain PAPER products."""
    out = tmp_path / "seg.xlsx"
    ExcelTemplateGenerator().generate(
        customers=["Metro Tyres"],
        products_by_segment={
            "CARPET": ["FGCB200-200", "FGCB200-220"],
            "PAPER": ["FGC8200-200", "FGP8040-200"],
            "CONSTRUCTION B2B IH": ["FGABSSHS400-1"],
        },
        output_path=out,
    )
    wb = load_workbook(out)
    lists = wb["_lists"]
    carpet_col = None
    for col in range(5, 20):
        if lists.cell(1, col).value == "CARPET":
            carpet_col = col
            break
    assert carpet_col is not None
    carpet_codes = []
    r = 2
    while lists.cell(r, carpet_col).value:
        carpet_codes.append(lists.cell(r, carpet_col).value)
        r += 1
    assert carpet_codes == ["FGCB200-200", "FGCB200-220"]
    assert "FGC8200-200" not in carpet_codes
    assert "SEG_CARPET" in wb.defined_names
    assert "SEG_CONSTRUCTION_B2B_IH" in wb.defined_names
    sheet = wb["Sales Report"]
    # Conditional formatting flags stale Product after Segment change (no VBA)
    assert len(sheet.conditional_formatting._cf_rules) >= 1
    formulas = [str(dv.formula1) for dv in sheet.data_validations.dataValidation]
    assert any("INDIRECT" in f for f in formulas)


def test_segment_range_name_sanitizes_spaces():
    from app.utils.excel_names import segment_range_name

    assert segment_range_name("CARPET") == "SEG_CARPET"
    assert segment_range_name("CONSTRUCTION B2B IH") == "SEG_CONSTRUCTION_B2B_IH"
    assert segment_range_name("123Start").startswith("SEG_N_")


@pytest.mark.asyncio
async def test_generate_requires_masters(db: Session):
    # Clear via replace with empty is not allowed by parser; soft-delete manually
    CustomerMasterRepository(db).soft_delete_all_active()
    ProductMasterRepository(db).soft_delete_all_active()
    svc = TemplateGenerationService(db)
    with pytest.raises(ValidationAppError):
        svc.generate(actor="test")


@pytest.mark.asyncio
async def test_bulk_replace_5000_customers_efficient(db: Session, tmp_path: Path):
    """Scale smoke: 5000 customers replace via bulk_insert_mappings."""
    names = [f"Bulk Customer {i:05d}" for i in range(5000)]
    path = _write_customer_master(tmp_path / "c5k.xlsx", names)
    svc = CustomerMasterService(db)
    result = await svc.upload_and_replace(_Upload(path), actor="test")
    assert result.success
    assert result.records_imported == 5000
    assert CustomerMasterRepository(db).count() == 5000
    assert result.processing_time_ms > 0


@pytest.mark.asyncio
async def test_advisory_lock_serializes_replace(db: Session):
    """Advisory lock helper is callable within the session transaction."""
    from app.utils.db_locks import CUSTOMER_MASTER_REPLACE_LOCK, acquire_xact_lock

    acquire_xact_lock(db, CUSTOMER_MASTER_REPLACE_LOCK, label="test")
    CustomerMasterRepository(db).replace_all(["Lock Test A", "Lock Test B"])
    assert set(CustomerMasterRepository(db).list_names()) == {"Lock Test A", "Lock Test B"}
