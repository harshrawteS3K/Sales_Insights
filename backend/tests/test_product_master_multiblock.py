"""Product Master multi-block + header-alias regression tests."""

from pathlib import Path

import pytest
from openpyxl import Workbook

from app.exceptions import ExcelProcessingError
from app.integrations.excel.headers import compact_header, match_product_master_field
from app.integrations.excel.parser import ExcelParserService
from app.integrations.excel.product_master_parser import parse_product_master_workbook


def _write_single_table(path: Path, rows: list[tuple[str, str]], *, headers=None) -> Path:
    wb = Workbook()
    ws = wb.active
    hdr = headers or ("Industry Type Description", "Product Code")
    ws.append(list(hdr))
    for industry, code in rows:
        ws.append([industry, code])
    wb.save(path)
    return path


def _write_four_blocks(path: Path) -> Path:
    """APCOTEX-style side-by-side Product Master (B:C, E:F, H:I, K:L)."""
    wb = Workbook()
    ws = wb.active
    # Row 1 headers at columns B,C / E,F / H,I / K,L (1-based: 2,3 / 5,6 / 8,9 / 11,12)
    blocks = [
        (2, "PAPER", [f"FGP{i}-200" for i in range(1, 6)]),
        (5, "CARPET", [f"FGCB{i}-200" for i in range(1, 6)]),
        (8, "CONSTRUCTION B2B", [f"FGB{i}-1" for i in range(1, 6)]),
        (11, "CONSTRUCTION B2C", [f"FGC{i}-1" for i in range(1, 6)]),
    ]
    for col, _industry, _codes in blocks:
        ws.cell(1, col, "Industry Type Description")
        ws.cell(1, col + 1, "Product Code")
    for col, industry, codes in blocks:
        for i, code in enumerate(codes, start=2):
            ws.cell(i, col, industry)
            ws.cell(i, col + 1, code)
    # Noise / blank column D, G, J
    ws.cell(1, 1, "")  # column A empty
    wb.save(path)
    return path


def test_scenario1_single_table_imports(tmp_path: Path):
    path = _write_single_table(
        tmp_path / "single.xlsx",
        [("PAPER", "FGP100-200"), ("CARPET", "FGCB200-200")],
    )
    records, dups = ExcelParserService().parse_product_master(path)
    assert len(records) == 2
    assert dups == 0
    assert {r["product_code"] for r in records} == {"FGP100-200", "FGCB200-200"}


def test_scenario2_four_side_by_side_blocks(tmp_path: Path):
    path = _write_four_blocks(tmp_path / "multi.xlsx")
    result = parse_product_master_workbook(path)
    assert len(result.blocks) == 4
    assert len(result.records) == 20
    assert result.duplicates_ignored == 0
    industries = {r["industry_type"] for r in result.records}
    assert industries == {"PAPER", "CARPET", "CONSTRUCTION B2B", "CONSTRUCTION B2C"}


def test_scenario3_mixed_spacing_headers(tmp_path: Path):
    path = tmp_path / "spaced.xlsx"
    _write_single_table(
        path,
        [("PAPER", "FGP1")],
        headers=("  Industry   Type  Description  ", "Product_Code"),
    )
    records, _ = ExcelParserService().parse_product_master(path)
    assert len(records) == 1
    assert records[0]["product_code"] == "FGP1"


def test_scenario4_industry_type_description_alias():
    assert match_product_master_field("Industry Type Description") == "industry_type"
    assert match_product_master_field("industry_type_description") == "industry_type"
    assert compact_header("Industry Type Description") == "industrytypedescription"


def test_scenario5_product_code_alias():
    assert match_product_master_field("Product Code") == "product_code"
    assert match_product_master_field("PRODUCT CODE") == "product_code"
    assert match_product_master_field("ProductCode") == "product_code"


def test_scenario6_blank_rows_ignored(tmp_path: Path):
    path = tmp_path / "blanks.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Industry Type Description", "Product Code"])
    ws.append(["PAPER", "A1"])
    ws.append([None, None])
    ws.append(["", ""])
    ws.append(["PAPER", "A2"])
    wb.save(path)
    records, _ = ExcelParserService().parse_product_master(path)
    assert [r["product_code"] for r in records] == ["A1", "A2"]


def test_scenario7_duplicate_product_codes_deduped(tmp_path: Path):
    path = _write_single_table(
        tmp_path / "dup.xlsx",
        [("PAPER", "SAME-1"), ("CARPET", "SAME-1"), ("PAPER", "SAME-1")],
    )
    records, dups = ExcelParserService().parse_product_master(path)
    assert len(records) == 1
    assert dups == 2
    assert records[0]["product_code"] == "SAME-1"


def test_scenario8_invalid_block_skipped_valid_imported(tmp_path: Path):
    """One incomplete block (header only) must not block a valid neighboring block."""
    path = tmp_path / "mixed_blocks.xlsx"
    wb = Workbook()
    ws = wb.active
    # Valid block B:C
    ws.cell(1, 2, "Industry Type Description")
    ws.cell(1, 3, "Product Code")
    ws.cell(2, 2, "PAPER")
    ws.cell(2, 3, "FGP100")
    # Invalid / empty block E:F — headers only, no data
    ws.cell(1, 5, "Industry Type Description")
    ws.cell(1, 6, "Product Code")
    # Valid block H:I
    ws.cell(1, 8, "Industry Type Description")
    ws.cell(1, 9, "Product Code")
    ws.cell(2, 8, "CARPET")
    ws.cell(2, 9, "FGCB200")
    wb.save(path)

    result = parse_product_master_workbook(path)
    assert len(result.blocks) == 3
    codes = {r["product_code"] for r in result.records}
    assert codes == {"FGP100", "FGCB200"}


def test_missing_headers_still_errors(tmp_path: Path):
    path = tmp_path / "bad.xlsx"
    wb = Workbook()
    wb.active.append(["Foo", "Bar"])
    wb.active.append(["x", "y"])
    wb.save(path)
    with pytest.raises(ExcelProcessingError) as exc:
        parse_product_master_workbook(path)
    assert "missing required columns" in str(exc.value).lower() or "could not find" in str(exc.value).lower()
