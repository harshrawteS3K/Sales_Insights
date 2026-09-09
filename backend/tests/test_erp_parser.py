"""ERP Excel parser foundation tests."""

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from app.erp_parser import ERPParserService
from app.erp_parser.header_mapper import best_field_match, map_headers
from app.erp_parser.row_extractor import extract_rows
from app.erp_parser.workbook_detector import list_candidate_sheets


def _save(wb: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def test_single_sheet_erp(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Party Name", "Material", "Dispatch Qty"])
    ws.append(["AKS RUGS", "P1", 10])
    ws.append(["Beta Corp", "P2", 25.5])
    path = _save(wb, tmp_path / "single.xlsx")

    result = ERPParserService().parse_workbook(path)
    assert result.sheet_name == "Sales"
    assert result.imported_rows == 2
    assert result.rows[0]["customer_name"] == "AKS RUGS"
    assert result.rows[0]["product"] == "P1"
    assert float(result.rows[0]["sales_quantity"]) == 0.01  # 10 KG → MT
    assert result.parsed_rows[0].unit == "MT"
    assert result.overall_confidence >= 75
    assert result.parsed_rows[0].segment == ""
    assert result.parsed_rows[0].distributor == ""


def test_multi_sheet_picks_sales_sheet(tmp_path: Path):
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = "Annual Sales Statement 2026"
    cover["A2"] = "Confidential"

    junk = wb.create_sheet("Notes")
    junk["A1"] = "Please see Sales tab"

    sales = wb.create_sheet("Dispatch")
    sales.append(["Customer", "Item Code", "Qty"])
    sales.append(["Cust A", "FG-1", 100])
    sales.append(["Cust B", "FG-2", 200])
    path = _save(wb, tmp_path / "multi.xlsx")

    result = ERPParserService().parse_workbook(path)
    assert result.sheet_name == "Dispatch"
    assert result.imported_rows == 2
    assert "Dispatch" in result.candidate_sheets


def test_blank_sheets_ignored(tmp_path: Path):
    wb = Workbook()
    empty = wb.active
    empty.title = "Empty"
    sales = wb.create_sheet("Data")
    sales.append(["Buyer", "Product", "Quantity"])
    sales.append(["X", "Y", 1])
    path = _save(wb, tmp_path / "blank.xlsx")

    candidates = list_candidate_sheets(path)
    assert "Empty" not in candidates
    assert "Data" in candidates


def test_merged_header_workbook(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "ERP"
    ws.merge_cells("A1:C1")
    ws["A1"] = "Monthly Dispatch Report — Plant 1"
    ws.append([])  # row 2 blank after merge write — use explicit rows
    ws["A3"] = "Party Name"
    ws["B3"] = "Material Description"
    ws["C3"] = "Net Qty"
    ws["A4"] = "Omega Mills"
    ws["B4"] = "Latex-100"
    ws["C4"] = 55
    path = _save(wb, tmp_path / "merged.xlsx")

    result = ERPParserService().parse_workbook(path)
    assert result.imported_rows == 1
    assert result.rows[0]["customer_name"] == "Omega Mills"
    assert result.rows[0]["product"] == "Latex-100"
    assert float(result.rows[0]["sales_quantity"]) == 0.055  # 55 KG → MT


def test_different_header_names_fuzzy(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Account Name", "FG Code", "Billed Qty"])
    ws.append(["Acme", "G7", 3])
    path = _save(wb, tmp_path / "aliases.xlsx")

    result = ERPParserService().parse_workbook(path)
    assert result.imported_rows == 1
    mapping_fields = {m["field"]: m for m in result.mapping}
    assert mapping_fields["customer"]["confidence"] >= 80
    assert mapping_fields["product"]["confidence"] >= 80
    assert mapping_fields["quantity"]["confidence"] >= 80


def test_totals_ignored(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Customer Name", "Product", "Sales Quantity"])
    ws.append(["A", "P1", 10])
    ws.append(["B", "P2", 20])
    ws.append(["TOTAL", "", 30])
    ws.append(["Grand Total", "x", 99])
    path = _save(wb, tmp_path / "totals.xlsx")

    result = ERPParserService().parse_workbook(path)
    assert result.imported_rows == 2
    names = {r["customer_name"] for r in result.rows}
    assert names == {"A", "B"}


def test_quantity_parsing(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Client", "Item", "Invoice Qty"])
    ws.append(["C1", "I1", "1,250.5"])
    ws.append(["C2", "I2", 10])
    path = _save(wb, tmp_path / "qty.xlsx")

    result = ERPParserService().parse_workbook(path)
    assert result.imported_rows == 2
    assert float(result.rows[0]["sales_quantity"]) == 1.2505  # 1,250.5 KG → MT
    assert isinstance(result.parsed_rows[0].quantity, Decimal)


def test_monthly_pivot_erp(tmp_path: Path):
    """SOUTH NBR-style: months → FY quarterly rows (not one annual total)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(
        [
            "Item Description",
            "Customer Code",
            "Customer Name",
            "APRIL",
            "MAY",
            "JUNE",
            "JULY",
            "AUGUST",
            "SEPTEMBER",
        ]
    )
    ws.append(["NBR - APCOFLEX N 285", "TN-1", "A G RUBBER", None, None, None, 70, None, None])
    ws.append(["NBR - APCOFLEX N 745", "TN-1", "A G RUBBER", 525, 1015, 1015, 1995, 525, None])
    path = _save(wb, tmp_path / "monthly_pivot.xlsx")

    result = ERPParserService().parse_workbook(path, fiscal_year_start=2025)
    assert result.imported_rows == 3  # one Q2 row for N285 + Q1+Q2 for N745
    periods = {r.get("period") for r in result.rows}
    assert periods == {"Q1 2025", "Q2 2025"}
    q1 = next(
        r
        for r in result.rows
        if r["product"] == "NBR - APCOFLEX N 745" and r["period"] == "Q1 2025"
    )
    assert float(q1["sales_quantity"]) == (525 + 1015 + 1015) / 1000
    q2_n285 = next(r for r in result.rows if r["product"] == "NBR - APCOFLEX N 285")
    assert q2_n285["period"] == "Q2 2025"
    assert float(q2_n285["sales_quantity"]) == 0.07  # 70 KG → MT
    assert result.overall_confidence >= 75
    qty_map = next(m for m in result.mapping if m["field"] == "quantity")
    assert qty_map["method"] == "monthly_sum"


def test_header_mapper_priority():
    label, score, method = best_field_match("party name", "customer")
    assert label == "Customer Name"
    assert score >= 90
    assert method in {"exact", "synonym", "fuzzy"}

    mapped = map_headers(["Sr. No.", "Party Name", "Material", "Dispatch Qty"])
    assert mapped["positions"]["customer"] == 1
    assert mapped["positions"]["product"] == 2
    assert mapped["positions"]["quantity"] == 3


def test_preview_payload_shape(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Consignee", "Grade", "Qty"])
    ws.append(["Z", "G1", 7])
    path = _save(wb, tmp_path / "preview.xlsx")

    preview = ERPParserService().preview(path)
    assert "sheet_name" in preview
    assert "confidence" in preview
    assert "mapping" in preview
    assert preview["row_count"] == 1
    assert preview["rows"][0]["product"] == "G1"


def test_extract_rows_helper():
    matrix = [
        ["Party Name", "Material", "Dispatch Qty"],
        ["A", "P", 1],
        ["SUB TOTAL", "", 1],
        ["", "", ""],
        ["B", "Q", 2],
    ]
    out = extract_rows(
        matrix,
        header_row_index=0,
        positions={"customer": 0, "product": 1, "quantity": 2},
    )
    assert len(out["rows"]) == 2
    assert out["skipped_total"] >= 1
