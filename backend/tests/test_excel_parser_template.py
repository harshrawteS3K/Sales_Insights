"""Regression tests for template-driven APCOTEX Excel parser (quarterly)."""

from pathlib import Path

import pytest
from openpyxl import Workbook

from app.exceptions import ExcelProcessingError
from app.integrations.excel.headers import normalize_header
from app.integrations.excel.parser import ExcelParserService
from tests.workbook_helpers import build_official_workbook


@pytest.fixture
def parser() -> ExcelParserService:
    return ExcelParserService()


def _save(wb: Workbook, path: Path) -> Path:
    wb.save(path)
    return path


def test_normalize_header_variants():
    assert normalize_header("Sr. No.") == "sr no"
    assert normalize_header("Customer Name") == "customer name"
    assert normalize_header("Name of Customer") == "name of customer"
    assert normalize_header("Product") == "product"
    assert normalize_header("Product (FG Code)") == "product fg code"
    assert normalize_header("Sales Quantity") == "sales quantity"
    assert normalize_header("Reporting Quarter") == "reporting quarter"


def test_official_quarterly_apcotex_template(parser: ExcelParserService, tmp_path: Path):
    path = build_official_workbook(
        tmp_path / "official.xlsx",
        distributor="Navneet Goel",
        reporting_month="Q3 2026",
        rows=[
            (1, "M/S AKS RUGS CO.", "Carpet", "APCOTEX CB 4600", 3300),
            (2, "M/S BETA CO.", "Carpet", "APCOTEX CB 300", 100),
        ],
    )
    result = parser.parse_sales_report(path)

    assert result.mapping_strategy == "official_template"
    assert result.template_name == "Official APCOTEX Template"
    assert result.imported_rows == 2
    assert result.incomplete_rows == 0
    assert result.quality_score >= 80
    assert result.reporting_month == "Q3 2026"
    assert result.distributor_details["name"] == "Navneet Goel"
    assert result.rows[0].period == "Q3 2026"
    assert result.rows[0].product == "APCOTEX CB 4600"
    assert result.rows[0].quantity == 3300


def test_legacy_month_labels_still_parse(parser: ExcelParserService, tmp_path: Path):
    path = build_official_workbook(
        tmp_path / "legacy.xlsx",
        distributor="Dist A",
        reporting_month="June 2026",
        include_stock=False,
        use_legacy_month_labels=True,
        rows=[(1, "Cust A", "Carpet", "CB 300", 10)],
    )
    result = parser.parse_sales_report(path)
    assert result.imported_rows == 1
    assert result.reporting_month == "June 2026"


def test_legacy_period_column_fallback(parser: ExcelParserService, tmp_path: Path):
    """Old Excel with Period column and no Reporting Quarter header still works."""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Name of Distributor"
    ws["B1"] = "Legacy Dist"
    ws["A2"] = "Company Name"
    ws["B2"] = "Legacy Dist"
    headers = ["Sr. No.", "Name of Customer", "Segment", "Product", "Quantity", "Period"]
    for i, h in enumerate(headers, 1):
        ws.cell(9, i, h)
    ws.cell(10, 1, 1)
    ws.cell(10, 2, "Cust")
    ws.cell(10, 3, "Carpet")
    ws.cell(10, 4, "P1")
    ws.cell(10, 5, 5)
    ws.cell(10, 6, "May 2026")
    path = _save(wb, tmp_path / "legacy_period.xlsx")
    result = parser.parse_sales_report(path)
    assert result.imported_rows == 1
    assert result.reporting_month == "May 2026"


def test_missing_required_column_fails(parser: ExcelParserService, tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Name of Person"
    ws["B1"] = "Someone"
    ws["A2"] = "Reporting Quarter"
    ws["B2"] = "Q1 2026"
    for i, h in enumerate(["Sr. No.", "Customer Name", "Segment", "Sales Quantity"], 1):
        ws.cell(6, i, h)
    ws.cell(7, 1, 1)
    ws.cell(7, 2, "Cust")
    ws.cell(7, 3, "Carpet")
    ws.cell(7, 4, 10)
    path = _save(wb, tmp_path / "missing_product.xlsx")
    with pytest.raises(ExcelProcessingError):
        parser.parse_sales_report(path)


def test_missing_reporting_quarter_fails(parser: ExcelParserService, tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Name of Person"
    ws["B1"] = "Someone"
    ws["A2"] = "Company Name"
    ws["B2"] = "Someone Co"
    for i, h in enumerate(
        ["Sr. No.", "Customer Name", "Segment", "Product", "Sales Quantity"], 1
    ):
        ws.cell(6, i, h)
    ws.cell(7, 1, 1)
    ws.cell(7, 2, "Cust")
    ws.cell(7, 3, "Carpet")
    ws.cell(7, 4, "P1")
    ws.cell(7, 5, 10)
    path = _save(wb, tmp_path / "no_quarter.xlsx")
    with pytest.raises(ExcelProcessingError) as exc:
        parser.parse_sales_report(path)
    assert "Reporting Quarter" in str(exc.value)
